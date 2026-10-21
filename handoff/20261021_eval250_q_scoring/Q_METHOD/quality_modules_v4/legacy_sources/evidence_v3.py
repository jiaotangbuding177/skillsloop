"""Read-only, method-blind evidence for the frozen 750 selected deliveries.

This module does not contact a model, rerun an Actor, calculate TSR or assign Q.
Native files are never edited. Missing files and failed evidence are distinct.
"""
from __future__ import annotations

import csv
import email
import hashlib
import io
import json
import math
import re
import shutil
import subprocess
import sys
import threading
import zipfile
from datetime import date, datetime
from email import policy
from pathlib import Path, PurePosixPath
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent.parent
RUNTIME = ROOT / "evidence_runtime_v3"
SOURCE_PACKETS = PROJECT / "output/20261017_eval250_rubrics/sources/task_packets.jsonl"
SCHEMA = "eval250-native-quality-evidence-v3"
MAX_FILE_BYTES = 128 * 1024 * 1024
MAX_ARCHIVE_BYTES = 256 * 1024 * 1024
PDF_DPI = 120
_LOCK = threading.Lock()
_PACKETS = None
_RENDER_SEMAPHORE = threading.BoundedSemaphore(2)
_SOURCE_RENDER_LOCKS = {}
TEXT_SUFFIXES = {".txt", ".md", ".json", ".html", ".htm", ".css", ".js", ".mjs", ".ts", ".py", ".svg", ".xml", ".yaml", ".yml", ".sh", ".sql", ".ini", ".log"}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def write(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf8")
    temporary.replace(path)


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return {"native_nonfinite": str(value)}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def safe_relative(value):
    p = PurePosixPath(str(value).replace("\\", "/"))
    if p.is_absolute() or not p.parts or any(x in {"..", "."} for x in p.parts) or ":" in str(p):
        raise ValueError("invalid_relative_evidence_path")
    return p


def _error(exc, role="evidence"):
    # Do not leak original attempt paths / method names through exception strings.
    return {"kind": type(exc).__name__, "role": role, "message": str(exc).split("\n")[0].replace(str(PROJECT), "[project]")}


def _archive(path):
    z = zipfile.ZipFile(path)
    if sum(x.file_size for x in z.infolist()) > MAX_ARCHIVE_BYTES:
        z.close()
        raise ValueError("archive_resource_bound_no_silent_truncation")
    return z


def _xml_text(element):
    return "".join(x.text or "" for x in element.iter() if x.tag.rsplit("}", 1)[-1] in {"t", "instrText", "delText"})


def extract_native(path, derived_dir=None, depth=0):
    """Complete parser output, including formulas and caches separately.

    A resource bound raises an explicit failure. It never samples or truncates.
    """
    path = Path(path)
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("file_resource_bound_no_silent_truncation")
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES or suffix in {".csv", ".tsv", ".ics"}:
        raw = path.read_bytes()
        encoding = "utf-8-sig"
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            encoding = "gb18030"
            text = raw.decode(encoding, errors="strict")
        result = {"parser": "strict-text", "encoding": encoding, "text": text,
                  "lines": [{"locator": f"line:{i}", "text": s} for i, s in enumerate(text.splitlines(), 1)],
                  "content_complete": True, "limitations": []}
        if suffix == ".json":
            result["parsed_json"] = json.loads(text)
        if suffix in {".csv", ".tsv"}:
            result["rows"] = list(csv.reader(io.StringIO(text, newline=""), delimiter="\t" if suffix == ".tsv" else ","))
        if suffix == ".ics":
            unfolded = re.sub(r"\r?\n[ \t]", "", text)
            result["unfolded_lines"] = unfolded.splitlines()
            result["calendar_structure"] = [{"line": i, "property": s.split(":", 1)[0], "value": s.split(":", 1)[1] if ":" in s else ""}
                                              for i, s in enumerate(unfolded.splitlines(), 1)]
            result["limitations"].append("Calendar source parsed; application import was not performed by this parser")
            try:
                from dateutil.rrule import rrulestr
                from zoneinfo import ZoneInfo
                from datetime import timedelta
                events, current = [], None
                for line in unfolded.splitlines():
                    if line == "BEGIN:VEVENT":
                        current = []
                    elif line == "END:VEVENT" and current is not None:
                        events.append(current)
                        current = None
                    elif current is not None and ":" in line:
                        k, v = line.split(":", 1)
                        current.append((k, v))
                def parse_dt(key, value):
                    fmt = "%Y%m%d" if len(value) == 8 else "%Y%m%dT%H%M%S" + ("Z" if value.endswith("Z") else "")
                    parsed = datetime.strptime(value, fmt)
                    tz = re.search(r"(?:^|;)TZID=([^;:]+)", key)
                    if value.endswith("Z"):
                        from datetime import timezone
                        parsed = parsed.replace(tzinfo=timezone.utc)
                    elif tz:
                        parsed = parsed.replace(tzinfo=ZoneInfo(tz.group(1).strip('"')))
                    return parsed
                instances, overrides, bases = [], {}, []
                for fields in events:
                    properties = {k.split(";", 1)[0]: (k, v) for k, v in fields}
                    uid = properties.get("UID", ("", ""))[1]
                    if "DTSTART" not in properties:
                        continue
                    start = parse_dt(*properties["DTSTART"])
                    end = parse_dt(*properties["DTEND"]) if "DTEND" in properties else start
                    event = {"uid": uid, "start": start.isoformat(), "end": end.isoformat(), "fields": fields}
                    if "RECURRENCE-ID" in properties:
                        recurrence = parse_dt(*properties["RECURRENCE-ID"]).isoformat()
                        overrides[(uid, recurrence)] = event
                        continue
                    dates = [start]
                    if "RRULE" in properties:
                        rule = properties["RRULE"][1]
                        if "COUNT=" not in rule and "UNTIL=" not in rule:
                            raise ValueError("unbounded_calendar_series_no_silent_sampling")
                        dates = list(rrulestr(rule, dtstart=start))
                        if len(dates) > 1000:
                            raise ValueError("calendar_series_resource_bound_no_silent_sampling")
                    excluded = set()
                    for key, value in fields:
                        if key.split(";", 1)[0] == "EXDATE":
                            excluded.update(parse_dt(key, x).isoformat() for x in value.split(","))
                    for dt in dates:
                        if dt.isoformat() not in excluded:
                            bases.append({"uid": uid, "start": dt.isoformat(), "end": (dt + (end-start)).isoformat(), "original_recurrence_start": dt.isoformat()})
                for instance in bases:
                    override = overrides.pop((instance["uid"], instance["original_recurrence_start"]), None)
                    instances.append({**instance, **override} if override else instance)
                instances.extend(overrides.values())
                result["event_components"] = events
                result["actual_recurrence_expansion"] = {"parser": "datetime/zoneinfo/dateutil.rrule", "instances": instances,
                    "scope": "Actual local standards-based recurrence expansion; not a calendar application import"}
            except Exception as exc:
                result["recurrence_expansion_failure"] = type(exc).__name__ + ":" + str(exc)
        return result
    if suffix == ".eml":
        message = email.message_from_bytes(path.read_bytes(), policy=policy.default)
        parts = []
        for i, part in enumerate(message.walk(), 1):
            if part.is_multipart():
                continue
            payload = part.get_payload(decode=True) or b""
            item = {"locator": f"mime-part:{i}", "content_type": part.get_content_type(), "filename": part.get_filename(),
                    "bytes": len(payload), "sha256": hashlib.sha256(payload).hexdigest()}
            if part.get_content_maintype() == "text":
                item["text"] = payload.decode(part.get_content_charset() or "utf8", errors="strict")
            if part.get_filename():
                attachment_root = Path(derived_dir or path.parent / "mime-derived")
                attachment = attachment_root / (f"part-{i:03}-" + Path(part.get_filename()).name)
                attachment.parent.mkdir(parents=True, exist_ok=True)
                attachment.write_bytes(payload)
                try:
                    item["attachment_extraction"] = extract_native(attachment, attachment_root / f"part-{i:03}-derived", depth + 1)
                except Exception as exc:
                    item["attachment_extraction_failure"] = _error(exc, "mime-attachment")
            parts.append(item)
        return {"parser": "email.policy.default", "headers": list(message.raw_items()), "parts": parts, "content_complete": True,
                "limitations": ["MIME delivery draft only; email not sent or imported"]}
    if suffix in {".xlsx", ".xlsm"}:
        import openpyxl
        with _archive(path) as z:
            xml_parts = {n: z.read(n).decode("utf8") for n in z.namelist()
                         if re.match(r"xl/(?:charts|drawings|comments|externalLinks)/.*\.xml$", n)}
        formulas = openpyxl.load_workbook(path, data_only=False, keep_links=True)
        cache = openpyxl.load_workbook(path, data_only=True, keep_links=True)
        sheets = []
        try:
            for sheet in formulas:
                cells = []
                for cell in sorted(sheet._cells.values(), key=lambda c: (c.row, c.column)):
                    if cell.value is None:
                        continue
                    cells.append({"locator": f"sheet:{sheet.title}!{cell.coordinate}", "cell": cell.coordinate,
                                  "value": scalar(cell.value), "data_type": cell.data_type,
                                  "original_cached_value": scalar(cache[sheet.title][cell.coordinate].value),
                                  "cached_value_reliability": "NOT_RECALCULATED_MAY_BE_STALE",
                                  "number_format": cell.number_format, "style_id": cell.style_id,
                                  "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                                  "comment": cell.comment.text if cell.comment else None})
                sheets.append({"sheet": sheet.title, "state": sheet.sheet_state, "rows": sheet.max_row, "columns": sheet.max_column,
                               "cells": cells, "merged_ranges": [str(x) for x in sheet.merged_cells.ranges],
                               "freeze_panes": str(sheet.freeze_panes), "auto_filter": sheet.auto_filter.ref,
                               "print_area": str(sheet.print_area), "print_title_rows": sheet.print_title_rows,
                               "page_setup": str(sheet.page_setup), "page_margins": str(sheet.page_margins),
                               "row_dimensions": {str(k): {"height": v.height, "hidden": v.hidden} for k, v in sheet.row_dimensions.items()},
                               "column_dimensions": {str(k): {"width": v.width, "hidden": v.hidden} for k, v in sheet.column_dimensions.items()},
                               "data_validations": str(sheet.data_validations), "tables": list(sheet.tables)})
            return {"parser": f"openpyxl/{openpyxl.__version__}", "sheets": sheets,
                    "styles": [{"style_id": i, "font": str(formulas._fonts[s.fontId]), "fill": str(formulas._fills[s.fillId]),
                                "alignment": str(formulas._alignments[s.alignmentId]), "border": str(formulas._borders[s.borderId])}
                               for i, s in enumerate(formulas._cell_styles)],
                    "defined_names": [str(x) for x in formulas.defined_names.values()], "embedded_xml": xml_parts,
                    "calculation_properties": str(formulas.calculation), "content_complete": True,
                    "limitations": ["Original cached values are not evidence of current formula correctness; no original file was recalculated"]}
        finally:
            formulas.close()
            cache.close()
    if suffix in {".docx", ".pptx"}:
        with _archive(path) as z:
            names = z.namelist()
            if suffix == ".docx":
                selected = [n for n in names if re.match(r"word/(?:document|header\d*|footer\d*|footnotes|endnotes|comments|styles|numbering|settings)\.xml$", n)
                            or re.match(r"docProps/.*\.xml$", n) or re.match(r"customXml/.*\.xml$", n)]
            else:
                selected = [n for n in names if (n.startswith("ppt/") or n.startswith("docProps/")) and n.endswith(".xml")]
            parts = []
            for n in selected:
                xml = z.read(n).decode("utf8")
                tree = ET.fromstring(xml)
                paragraphs = []
                for index, p in enumerate((x for x in tree.iter() if x.tag.rsplit("}", 1)[-1] == "p"), 1):
                    paragraphs.append({"locator": f"part:{n}/paragraph:{index}", "text": _xml_text(p)})
                parts.append({"locator": f"part:{n}", "part": n, "paragraphs": paragraphs, "xml": xml})
            media = [{"member": n, "bytes": len(z.read(n)), "sha256": hashlib.sha256(z.read(n)).hexdigest()} for n in names if "/media/" in n and not n.endswith("/")]
            relationships = {n: z.read(n).decode("utf8") for n in names if n.endswith(".rels")}
        font_facts = None
        if suffix == ".pptx":
            font_facts = _pptx_font_facts(path)
        return {"parser": "complete-selected-OOXML-parts-v3", "parts": parts, "embedded_media": media,
                "font_facts": font_facts,
                "relationships": relationships, "content_complete": True,
                "limitations": ["OOXML structure and text are not rendered page appearance; animations/macros not executed"]}
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = [{"locator": f"page:{i}", "page": i, "text": p.extract_text() or "",
                  "width_pt": float(p.mediabox.width), "height_pt": float(p.mediabox.height), "rotation": p.get("/Rotate", 0),
                  "annotations": _pdf_safe_annotations(p),
                  "font_resources": _pdf_font_resources(p)}
                 for i, p in enumerate(reader.pages, 1)]
        return {"parser": "pypdf", "page_count": len(pages), "pages": pages,
                "fields": {k: _pdf_safe(v) for k, v in (reader.get_fields() or {}).items()},
                "metadata": {str(k): _pdf_safe(v) for k, v in (reader.metadata or {}).items()},
                "embedded_font_facts": _pdf_font_facts(reader),
                "content_complete": True, "limitations": ["Text extraction does not replace actual rendered pages"]}
    if suffix == ".zip":
        if depth > 2:
            raise ValueError("nested_archive_depth_bound_no_silent_truncation")
        members = []
        target = Path(derived_dir or path.parent / "zip-derived")
        with _archive(path) as z:
            for member in z.infolist():
                relative = safe_relative(member.filename)
                if member.is_dir():
                    continue
                data = z.read(member)
                p = target.joinpath(*relative.parts)
                p.parent.mkdir(parents=True, exist_ok=True)
                p.write_bytes(data)
                item = {"locator": f"zip-member:{member.filename}", "member": member.filename, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest()}
                try:
                    item["extraction"] = extract_native(p, p.parent / (p.name + ".derived"), depth + 1)
                except Exception as exc:
                    item["extraction_failure"] = _error(exc, "zip-member")
                members.append(item)
        return {"parser": "zipfile-full-members", "members": members, "content_complete": not any("extraction_failure" in x for x in members),
                "limitations": ["Archive inspected only; included code was not executed"]}
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}:
        from PIL import Image
        with Image.open(path) as im:
            im.verify()
        with Image.open(path) as im:
            return {"parser": "Pillow", "format": im.format, "width": im.width, "height": im.height,
                    "frames": getattr(im, "n_frames", 1), "content_complete": True,
                    "limitations": ["Image pixels supplied separately; metadata alone does not establish visual quality"]}
    return {"parser": "unsupported", "content_complete": False, "limitations": ["No trusted parser for this native format"]}


def _environment():
    p = RUNTIME / "environment.json"
    return json.loads(p.read_text(encoding="utf8")) if p.exists() else {}


def _pptx_font_facts(path):
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    facts = {"kind": "ORIGINAL_PPTX_FONT_DECLARATIONS_NOT_RENDERING_PROOF", "themes": [], "slide_font_properties": [],
             "compatibility_status": "UNKNOWN_CROSS_APPLICATION_FONT_RESOLUTION",
             "interpretation": "Latin, East Asian and complex-script declarations are separate. Theme inheritance and local fallback differ by application; fixed-LibreOffice fallback alone is not proof of an original PPTX business defect."}
    with _archive(path) as z:
        for name in z.namelist():
            if re.fullmatch(r"ppt/theme/theme\d+\.xml", name):
                tree = ET.fromstring(z.read(name))
                for group in ["majorFont", "minorFont"]:
                    node = tree.find(".//a:" + group, ns)
                    if node is not None:
                        facts["themes"].append({"part": name, "font_group": group,
                            "declarations": [{"slot": child.tag.rsplit("}", 1)[-1], **child.attrib} for child in node]})
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                tree = ET.fromstring(z.read(name))
                for index, prop in enumerate(tree.findall(".//a:rPr", ns) + tree.findall(".//a:defRPr", ns) + tree.findall(".//a:endParaRPr", ns), 1):
                    fonts = [{"slot": child.tag.rsplit("}", 1)[-1], **child.attrib} for child in prop if child.tag.rsplit("}", 1)[-1] in {"latin", "ea", "cs", "sym"}]
                    if fonts or prop.get("lang") or prop.get("altLang"):
                        facts["slide_font_properties"].append({"part": name, "font_property_index": index, "language": prop.get("lang"), "alternate_language": prop.get("altLang"), "fonts": fonts})
    return facts


def _pdf_safe(value, depth=0):
    """Serialize pypdf values without object reprs or memory addresses."""
    if depth > 12:
        return {"depth_limit": True}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "idnum") and hasattr(value, "generation"):
        return {"indirect_ref": [int(value.idnum), int(value.generation)]}
    if isinstance(value, dict):
        return {str(key): _pdf_safe(child, depth + 1)
                for key, child in sorted(value.items(), key=lambda item: str(item[0]))}
    if isinstance(value, (list, tuple)):
        return [_pdf_safe(child, depth + 1) for child in value]
    return str(value)


def _pdf_resolved(value):
    return value.get_object() if hasattr(value, "get_object") else value


def _pdf_font_resources(page):
    resources = _pdf_resolved(page.get("/Resources", {}))
    fonts = _pdf_resolved(resources.get("/Font", {})) if isinstance(resources, dict) else {}
    result = {}
    if not isinstance(fonts, dict):
        return result
    for key, font in sorted(fonts.items(), key=lambda item: str(item[0])):
        resolved = _pdf_resolved(font)
        result[str(key)] = {
            "object_ref": _pdf_safe(font),
            "base_font": _pdf_safe(resolved.get("/BaseFont", "")) if isinstance(resolved, dict) else None,
            "subtype": _pdf_safe(resolved.get("/Subtype", "")) if isinstance(resolved, dict) else None,
            "encoding": _pdf_safe(resolved.get("/Encoding", "")) if isinstance(resolved, dict) else None}
    return result


def _pdf_safe_annotations(page):
    annotations = _pdf_resolved(page.get("/Annots", []))
    if not isinstance(annotations, (list, tuple)):
        return _pdf_safe(annotations)
    return [_pdf_safe(_pdf_resolved(annotation)) for annotation in annotations]


def _pdf_font_facts(reader):
    facts = []
    def describe(font):
        font = font.get_object()
        descriptor = font.get("/FontDescriptor")
        descriptor = descriptor.get_object() if descriptor else {}
        return {"base_font": _pdf_safe(font.get("/BaseFont", "")), "subtype": _pdf_safe(font.get("/Subtype", "")),
                "encoding": _pdf_safe(font.get("/Encoding", "")), "descriptor_font_name": _pdf_safe(descriptor.get("/FontName", "")),
                "embedded_font_stream_present": any(key in descriptor for key in ["/FontFile", "/FontFile2", "/FontFile3"]),
                "descendant_fonts": [describe(x) for x in font.get("/DescendantFonts", [])]}
    for number, page in enumerate(reader.pages, 1):
        resources = page.get("/Resources")
        resources = resources.get_object() if resources else {}
        fonts = resources.get("/Font")
        fonts = fonts.get_object() if fonts else {}
        for key, font in fonts.items():
            facts.append({"page": number, "resource": str(key), **describe(font),
                          "interpretation": "Font objects from original PDF bytes. Font family, encoding and embedded stream are native facts, not a visual guess."})
    return facts


def _sheet_wide_copy(original, destination):
    """Change only print-view layout in a derivative, never data/style/width."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    tag = lambda n: "{" + main + "}" + n
    with _archive(original) as z:
        members = {x.filename: (x, z.read(x.filename)) for x in z.infolist()}
    changed = {}
    invariant_sheets = []
    for name, (_, raw) in members.items():
        if not re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
            continue
        tree = ET.fromstring(raw)
        before = ET.fromstring(raw)
        sheetpr = tree.find(tag("sheetPr"))
        if sheetpr is None:
            sheetpr = ET.Element(tag("sheetPr"))
            tree.insert(0, sheetpr)
        setuppr = sheetpr.find(tag("pageSetUpPr"))
        if setuppr is None:
            setuppr = ET.SubElement(sheetpr, tag("pageSetUpPr"))
        setuppr.set("fitToPage", "1")
        setuppr.set("autoPageBreaks", "0")
        setup = tree.find(tag("pageSetup"))
        if setup is None:
            setup = ET.SubElement(tree, tag("pageSetup"))
        setup.set("paperSize", "8")
        setup.set("orientation", "landscape")
        setup.set("fitToWidth", "1")
        setup.set("fitToHeight", "0")
        setup.attrib.pop("scale", None)
        def content_signature(node):
            return {"sheetData": ET.tostring(node.find(tag("sheetData")), encoding="unicode"),
                    "cols": ET.tostring(node.find(tag("cols")), encoding="unicode") if node.find(tag("cols")) is not None else None,
                    "dimension": dict(node.find(tag("dimension")).attrib) if node.find(tag("dimension")) is not None else None,
                    "mergeCells": ET.tostring(node.find(tag("mergeCells")), encoding="unicode") if node.find(tag("mergeCells")) is not None else None}
        left, right = content_signature(before), content_signature(tree)
        if left != right:
            raise ValueError("sheet_wide_copy_changed_data_styles_or_column_widths")
        changed[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
        invariant_sheets.append({"part": name, "cell_formula_style_width_signature_sha256": hashlib.sha256(canonical(left).encode()).hexdigest(),
                                 "data_formula_style_column_widths_identical": True})
    wb = ET.fromstring(members["xl/workbook.xml"][1])
    rels = ET.fromstring(members["xl/_rels/workbook.xml.rels"][1])
    target_by_id = {r.get("Id"): r.get("Target") for r in rels}
    names = wb.find(tag("definedNames"))
    if names is None:
        names = ET.Element(tag("definedNames"))
        wb.append(names)
    for index, sheet in enumerate(wb.find(tag("sheets"))):
        target = target_by_id[sheet.get("{" + rel + "}id")]
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        if part not in members:
            continue
        tree = ET.fromstring(members[part][1])
        positions = [c.get("r") for c in tree.findall(".//" + tag("c")) if c.get("r") and any(c.find(tag(n)) is not None for n in ["v", "is", "f"])]
        if not positions:
            continue
        from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
        coords = [(column_index_from_string(coordinate_from_string(p)[0]), coordinate_from_string(p)[1]) for p in positions]
        lo_col, hi_col = min(x[0] for x in coords), max(x[0] for x in coords)
        lo_row, hi_row = min(x[1] for x in coords), max(x[1] for x in coords)
        for named in list(names):
            if named.get("name") == "_xlnm.Print_Area" and named.get("localSheetId") == str(index):
                names.remove(named)
        named = ET.SubElement(names, tag("definedName"), {"name": "_xlnm.Print_Area", "localSheetId": str(index)})
        quoted_sheet = "'" + sheet.get("name").replace("'", "''") + "'"
        named.text = f"{quoted_sheet}!${get_column_letter(lo_col)}${lo_row}:${get_column_letter(hi_col)}${hi_row}"
    def workbook_content_signature(raw):
        node = ET.fromstring(raw)
        defined = node.find(tag("definedNames"))
        if defined is not None:
            for child in list(defined):
                if child.get("name") == "_xlnm.Print_Area":
                    defined.remove(child)
            if not len(defined):
                node.remove(defined)
        return ET.tostring(node, encoding="unicode")
    changed["xl/workbook.xml"] = ET.tostring(wb, encoding="utf-8", xml_declaration=True)
    if workbook_content_signature(members["xl/workbook.xml"][1]) != workbook_content_signature(changed["xl/workbook.xml"]):
        raise ValueError("sheet_wide_copy_changed_workbook_content_outside_print_area")
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as out:
        for name, (info, raw) in members.items():
            out.writestr(info, changed.get(name, raw))
    with _archive(destination) as z:
        unchanged_parts = all(z.read(name) == raw for name, (_, raw) in members.items() if name not in changed)
    if not unchanged_parts:
        raise ValueError("sheet_wide_copy_changed_nonlayout_parts")
    return {"kind": "SHEET_WIDE_VIEWPORT", "changes": "Only pageSetup/pageSetUpPr plus print-area named ranges of all nonempty cells: A3 landscape, fit width=1,height=0. No data, formulas, styles, column widths or rows edited.",
            "unchanged_unmodified_package_parts": unchanged_parts, "workbook_content_outside_print_area_identical": True,
            "original_styles_xml_sha256": hashlib.sha256(members["xl/styles.xml"][1]).hexdigest() if "xl/styles.xml" in members else None,
            "sheet_invariance": invariant_sheets,
            "original_sha256": sha(original), "derived_sha256": sha(destination), "original_values_and_styles_preserved": True,
            "not_a_business_repair": True, "view_semantics": "Layout-only read viewport; judge original workbook readability here. Original default A4 pagination is retained separately as diagnosis, not a new XLSX print requirement."}


def _run(command, log_path, timeout=180):
    result = subprocess.run(command, capture_output=True, timeout=timeout)
    write(log_path, {"command": command, "returncode": result.returncode,
                     "stdout": result.stdout.decode("utf8", errors="replace"), "stderr": result.stderr.decode("utf8", errors="replace")})
    if result.returncode:
        raise RuntimeError(f"renderer_nonzero_exit_{result.returncode}")
    return result


def render_pdf(path, destination):
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    exe = _environment().get("pdftoppm_absolute_path") or shutil.which("pdftoppm")
    if not exe:
        raise RuntimeError("poppler_unavailable")
    prefix = destination / "page"
    with _RENDER_SEMAPHORE:
        _run([exe, "-r", str(PDF_DPI), "-png", str(path), str(prefix)], destination / "renderer_receipt.json", 180)
    pages = sorted(destination.glob("page-*.png"), key=lambda p: int(p.stem.rsplit("-", 1)[1]))
    from pypdf import PdfReader
    expected = len(PdfReader(path).pages)
    if len(pages) != expected:
        raise RuntimeError("render_page_count_mismatch_no_sampling")
    return [{"path": str(p.resolve()), "sha256": sha(p), "locator": f"page:{i}", "renderer": f"Poppler/{PDF_DPI}dpi", "page": i}
            for i, p in enumerate(pages, 1)]


def _render_office(path, destination):
    """Only a separately configured fixed runtime, derived files only."""
    environment = _environment()
    executable = environment.get("soffice_absolute_path")
    docker_image = environment.get("docker_image_id")
    if not executable and not docker_image:
        raise RuntimeError("native_office_renderer_unavailable_visual_items_unscored")
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    native = destination / "source"
    native.mkdir(exist_ok=True)
    copied = native / path.name
    shutil.copyfile(path, copied)
    converted = destination / "converted"
    converted.mkdir(exist_ok=True)
    profile = destination / "lo_profile"
    if executable:
        command = [executable, "-env:UserInstallation=" + profile.resolve().as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(converted), str(copied)]
    else:
        command = ["docker", "run", "--rm", "--network", "none", "--read-only", "--tmpfs", "/tmp:rw,nosuid,size=512m",
                   "--mount", f"type=bind,source={native.resolve()},target=/source,readonly",
                   "--mount", f"type=bind,source={converted.resolve()},target=/converted",
                   "--mount", f"type=bind,source={Path(environment.get('fixed_fonts_directory', 'C:/Windows/Fonts')).resolve()},target=/usr/share/fonts/windows,readonly",
                   docker_image, "libreoffice", "-env:UserInstallation=file:///tmp/lo_profile", "--headless", "--convert-to", "pdf", "--outdir", "/converted", "/source/" + path.name]
    with _RENDER_SEMAPHORE:
        _run(command, destination / "office_renderer_receipt.json", 180)
    pdfs = list(converted.glob("*.pdf"))
    if len(pdfs) != 1 or not pdfs[0].stat().st_size:
        raise RuntimeError("native_office_conversion_missing_pdf")
    pages = render_pdf(pdfs[0], destination / "pages")
    for page in pages:
        page["native_office_renderer"] = environment.get("soffice_version", docker_image or executable)
        page["derived_pdf_sha256"] = sha(pdfs[0])
        page["source_native_sha256"] = sha(path)
    return pages


def _recalculate_excel(path, destination):
    """Recalculate a derivative only, retain original formulas/cache separately."""
    import openpyxl
    destination = Path(destination)
    environment = _environment()
    image_id = environment.get("docker_image_id")
    executable = environment.get("soffice_absolute_path")
    if not image_id and not executable:
        raise RuntimeError("fixed_calculation_engine_unavailable")
    source = destination / "recalc_source"
    out = destination / "recalculated"
    source.mkdir(parents=True, exist_ok=True)
    out.mkdir(exist_ok=True)
    native = source / path.name
    shutil.copyfile(path, native)
    if image_id:
        command = ["docker", "run", "--rm", "--network", "none", "--read-only", "--tmpfs", "/tmp:rw,nosuid,size=512m",
                   "--mount", f"type=bind,source={source.resolve()},target=/source,readonly",
                   "--mount", f"type=bind,source={out.resolve()},target=/converted",
                   image_id, "libreoffice", "-env:UserInstallation=file:///tmp/recalc_profile", "--headless", "--convert-to", "xlsx", "--outdir", "/converted", "/source/" + path.name]
    else:
        command = [executable, "-env:UserInstallation=" + (destination / "recalc_profile").resolve().as_uri(), "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(native)]
    with _RENDER_SEMAPHORE:
        _run(command, destination / "calculation_receipt.private.json", 180)
    result = out / (path.stem + ".xlsx")
    if not result.is_file():
        raise RuntimeError("derived_recalculation_file_missing")
    original = openpyxl.load_workbook(path, data_only=False)
    derived = openpyxl.load_workbook(result, data_only=True)
    values = []
    try:
        for sheet in original:
            for cell in sheet._cells.values():
                if cell.data_type == "f":
                    values.append({"locator": f"sheet:{sheet.title}!{cell.coordinate}", "original_formula": scalar(cell.value),
                                   "derived_recalculated_value": scalar(derived[sheet.title][cell.coordinate].value),
                                   "derived_value_type": derived[sheet.title][cell.coordinate].data_type})
    finally:
        original.close()
        derived.close()
    return {"derived_file_sha256": sha(result), "engine": environment.get("soffice_version"), "formula_cells": values,
            "original_file_unchanged": sha(path) == sha(native),
            "scope": "Fixed-engine derivative cache only; engine incompatibility or unsupported formula is not automatically an Actor defect"}


def _source_pages(path, sha256):
    """Reuse source render bytes across all three arms without changing source."""
    environment_key = hashlib.sha256(canonical(_environment()).encode()).hexdigest()[:12]
    key = sha256 + "-" + environment_key
    with _LOCK:
        lock = _SOURCE_RENDER_LOCKS.setdefault(key, threading.Lock())
    cache = RUNTIME / "source_render_cache" / key
    with lock:
        index = cache / "render_index.json"
        if index.is_file():
            pages = json.loads(index.read_text(encoding="utf8"))
            if all(Path(x["path"]).is_file() and sha(x["path"]) == x["sha256"] for x in pages):
                return pages
            raise ValueError("source_render_cache_integrity_failure")
        if path.suffix.lower() == ".pdf":
            pages = render_pdf(path, cache)
        elif path.suffix.lower() in {".docx", ".xlsx", ".pptx"}:
            pages = _render_office(path, cache)
        else:
            return []
        write(index, pages)
        return pages


def _render_html(path, destination, bundle_root):
    environment = _environment()
    node, browser = environment.get("node_absolute_path"), environment.get("browser_absolute_path")
    script = RUNTIME / "browser_capture.cjs"
    if not node or not browser or not script.is_file():
        raise RuntimeError("fixed_browser_runtime_unavailable")
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    plan = {"root": str(Path(bundle_root).resolve()), "relative_path": str(path.relative_to(bundle_root)).replace("\\", "/"),
            "destination": str(destination.resolve()), "browser": browser,
            "playwright_module": environment["playwright_module_absolute_path"]}
    write(destination / "plan.private.json", plan)
    with _RENDER_SEMAPHORE:
        _run([node, str(script), str(destination / "plan.private.json")], destination / "browser_receipt.private.json", 360)
    result = json.loads((destination / "observations.json").read_text(encoding="utf8"))
    images = [{"path": str((destination / x["filename"]).resolve()), "sha256": sha(destination / x["filename"]),
               "locator": x["locator"], "renderer": "Chromium-fixed-viewports-v3", "viewport": x["viewport"],
               "filename": x["filename"], "state_ids": x.get("state_ids", []),
               "view_semantics": x.get("view_semantics", "REAL_INTERACTION_STATE_PIXELS")} for x in result["screenshots"]]
    return images, result


def _packet(task_id):
    global _PACKETS
    with _LOCK:
        if _PACKETS is None:
            _PACKETS = {r["task_id"]: r for r in [json.loads(x) for x in SOURCE_PACKETS.read_bytes().split(b"\n") if x.strip()]}
    return _PACKETS[task_id]


def _text(extraction):
    """Readable complete representation; structural data remains available too."""
    if extraction.get("text") is not None:
        return extraction["text"]
    return canonical(extraction)


def _evidence_atoms(evidence):
    """Index all actual scalar content without copying it into a second corpus.

    JSON pointers resolve against this complete evidence.json, not a projection.
    Plain strings quote literally; numbers/bools/null quote their JSON scalar.
    """
    atoms = []
    pointer_escape = lambda s: str(s).replace("~", "~0").replace("/", "~1")
    def add_record(record, pointer, role):
        eid = record.get("evidence_id")
        if not eid:
            return
        serial = 0
        def walk(value, location):
            nonlocal serial
            if isinstance(value, dict):
                for key, child in value.items():
                    if key in {"binding", "evidence_ids", "image_evidence_ids", "image_refs", "image_sha256"} or (key == "path" and location == pointer):
                        continue
                    if location == pointer and key in {"text", "frozen_units"} and (record.get("extraction") or record.get("observations") or record.get("arguments")):
                        continue  # exact parsed units below cover the same full content
                    if "/extraction/sheets/" in location and key in {"row_dimensions", "column_dimensions", "page_setup", "page_margins"}:
                        continue  # pixels and full structural evidence remain; no redundant metadata atom index
                    if re.search(r"/extraction/sheets/\d+/cells/\d+$", location) and key not in {"value", "original_cached_value", "number_format", "data_type", "hyperlink", "comment"}:
                        continue
                    if re.search(r"/extraction/parts/\d+/paragraphs/\d+$", location) and key != "text":
                        continue
                    if "/observations/" in location and key in {"viewports", "visible_controls", "screenshots"}:
                        continue  # state_catalog has unique complete text/target/menu facts, pictures have their own registry
                    if "/observations/" in location and key in {"bounds", "box", "scroll_x", "scroll_y", "scroll_width", "scroll_height", "viewport_width", "viewport_height", "viewport"}:
                        continue  # full geometry remains in original state; image metadata/claims handle geometric observations
                    walk(child, location + "/" + pointer_escape(key))
            elif isinstance(value, (list, tuple)):
                for index, child in enumerate(value):
                    walk(child, location + "/" + str(index))
            else:
                serial += 1
                atoms.append({"atom_id": f"{eid}:A{serial:05d}", "evidence_id": eid, "role": role,
                              "locator": record.get("locator", eid) + " @" + location,
                              "json_pointer": location, "value_type": "string" if isinstance(value, str) else "json_scalar"})
        walk(record, pointer)
    for container, role in [("sources", "input"), ("artifacts", "output"), ("functional_observations", "observation"), ("artifact_correspondence", "observation"), ("evidence_failures", "observation")]:
        for index, record in enumerate(evidence.get(container, [])):
            add_record(record, f"/{container}/{index}", role)
    return atoms


def prepare_job(task_rubric, inventory_row, directory):
    """Prepare a self-contained, arm-blind evidence.json for one selected job."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / "evidence.json"
    binding = {"task_id": task_rubric["task_id"], "blind_case_id": inventory_row["blind_case_id"],
               "rubric_sha256": hashlib.sha256(canonical(task_rubric).encode()).hexdigest(),
               "inventory_sha256": hashlib.sha256(canonical(inventory_row).encode()).hexdigest(),
               "evidence_source_sha256": sha(__file__), "environment_sha256": sha(RUNTIME / "environment.json") if (RUNTIME / "environment.json").is_file() else None,
               "support_files_sha256": {str(p.relative_to(ROOT)).replace("\\", "/"): sha(p) for p in [ROOT / "evidence_process.py", ROOT / "audit/scan_rubric_evidence.py", RUNTIME / "item_modality_contract.json", RUNTIME / "browser_capture.cjs"] if p.is_file()}}
    declared_modality = task_rubric.get("modality_contract_path")
    if declared_modality:
        binding["declared_modality_contract_sha256"] = sha(declared_modality)
    if target.exists():
        old = json.loads(target.read_text(encoding="utf8"))
        if old.get("binding") != binding:
            raise ValueError("existing_evidence_binding_mismatch_preserve_previous_use_fresh_destination")
        return old
    packet = _packet(task_rubric["task_id"])
    bundle = directory / "bundle"
    source_records, artifacts, visual, observations, failures = [], [], [], [], []
    missing, original_pins = [], []

    def fail(eid, locator, exc, kind="evidence"):
        # Keep failure messages strictly mechanical and path-free.
        failures.append({"evidence_id": eid, "locator": locator, "kind": kind,
                         "failure": type(exc).__name__, "reason": re.sub(r"[A-Za-z]:[\\/][^ ]+", "[private_path]", str(exc)).split("\n")[0]})

    for i, original in enumerate(packet["inputs"], 1):
        eid = f"I{i:03}"
        name = str(safe_relative(original["name"]))
        path = Path(original["path"])
        record = {"evidence_id": eid, "role": "input", "filename": name, "locator": f"input:{name}",
                  "sha256": original["sha256"], "frozen_units": original["units"]}
        try:
            if not path.is_file() or sha(path) != original["sha256"]:
                raise ValueError("source_missing_or_sha_changed")
            copied = bundle.joinpath(*safe_relative(name).parts)
            copied.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(path, copied)
            original_pins.append((path, original["sha256"]))
            record["extraction"] = extract_native(copied, directory / "source_members" / eid)
            record["text"] = _text(record["extraction"])
            record["parsing_status"] = "COMPLETE" if record["extraction"].get("content_complete") else "INCOMPLETE"
        except Exception as exc:
            record["parsing_status"] = "FAILED"
            fail(eid + ":failure", record["locator"], exc, "source")
        source_records.append(record)

    for record in source_records:
        if record.get("parsing_status") != "COMPLETE":
            continue
        name = record["filename"]
        copied = bundle.joinpath(*safe_relative(name).parts)
        if copied.suffix.lower() not in {".pdf", ".docx", ".xlsx", ".pptx"}:
            continue
        try:
            pages = _source_pages(copied, record["sha256"])
            for n, page in enumerate(pages, 1):
                visual.append({**page, "evidence_id": f"{record['evidence_id']}:V{n:03}", "role": "input",
                               "artifact": name, "relative_artifact": name, "locator": record["locator"] + "/" + page["locator"]})
            record["rendering_status"] = "COMPLETE"
        except Exception as exc:
            record["rendering_status"] = "FAILED"
            fail(record["evidence_id"] + ":render-failure", record["locator"], exc, "source-renderer")

    all_inventory = {x["relative_path"].replace("\\", "/"): x for x in inventory_row["all_output_artifacts"]}
    for required in inventory_row["required_artifacts"]:
        relative = required["relative_path"].replace("\\", "/")
        if relative not in all_inventory:
            all_inventory[relative] = {**required, "path": str(Path(inventory_row["attempt_ref"]) / "workspace" / relative)}
    required_names = {x["filename"] for x in inventory_row["required_artifacts"]}
    for i, (relative, entry) in enumerate(all_inventory.items(), 1):
        rel = safe_relative(relative)
        eid = f"O{i:03}"
        path = Path(entry["path"])
        record = {"evidence_id": eid, "artifact_id": eid, "filename": rel.name, "relative_path": str(rel),
                  "locator": f"output:{rel}", "required": rel.name in required_names, "evidence_ids": [eid],
                  "sha256": entry.get("sha256"), "presence": "PRESENT"}
        if not path.is_file() or path.stat().st_size == 0:
            record["presence"] = "MISSING_REQUIRED" if record["required"] else "MISSING_INVENTORIED_OUTPUT"
            record["text"] = "The required artifact is absent or zero bytes in the selected frozen delivery."
            record["parsing_status"] = "ABSENT"
            if record["required"]:
                missing.append({"evidence_id": eid, "locator": record["locator"], "filename": rel.name, "reason": record["presence"]})
            artifacts.append(record)
            continue
        try:
            actual = sha(path)
            if actual != entry.get("sha256"):
                raise ValueError("selected_output_sha_changed")
            copied = bundle.joinpath(*rel.parts)
            copied.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(path, copied)
            original_pins.append((path, actual))
            record["bytes"] = copied.stat().st_size
            record["extraction"] = extract_native(copied, directory / "output_members" / eid)
            record["text"] = _text(record["extraction"])
            record["parsing_status"] = "COMPLETE" if record["extraction"].get("content_complete") else "INCOMPLETE"
        except Exception as exc:
            record["parsing_status"] = "FAILED"
            fail(eid + ":parse-failure", record["locator"], exc, "parser")
        artifacts.append(record)

    # Render after all relative assets are copied, so local linked resources exist.
    for record in artifacts:
        if record["presence"] != "PRESENT" or record["parsing_status"] == "FAILED":
            continue
        eid, relative = record["evidence_id"], record["relative_path"]
        copied = bundle.joinpath(*safe_relative(relative).parts)
        extension = copied.suffix.lower()
        try:
            pages = []
            if extension == ".pdf":
                pages = render_pdf(copied, directory / "renders" / eid)
            elif extension in {".docx", ".pptx", ".xlsx", ".xlsm"}:
                pages = _render_office(copied, directory / "renders" / eid)
                if extension in {".xlsx", ".xlsm"}:
                    original_view = eid + ":original-print"
                    for p in pages:
                        p["view_id"] = original_view
                        p["view_semantics"] = "ORIGINAL_PRINT_DIAGNOSTIC_ONLY_NOT_XLSX_INTERACTIVE_READABILITY"
                    wide = directory / "renders" / eid / "sheet_wide_source" / copied.name
                    proof = _sheet_wide_copy(copied, wide)
                    wide_pages = _render_office(wide, directory / "renders" / eid / "sheet_wide_render")
                    wide_view = eid + ":sheet-wide"
                    for p in wide_pages:
                        p["view_id"] = wide_view
                        p["view_semantics"] = "SHEET_WIDE_VIEWPORT_LAYOUT_ONLY_DERIVATIVE_ORIGINAL_CONTENT_STYLE_WIDTHS_PRESERVED"
                        p["locator"] = "sheet-wide/" + p["locator"]
                    record["view_semantics"] = "Evaluate workbook contents and original-style whole-column readability using SHEET_WIDE_VIEWPORT. Original A4 pagination alone is not an XLSX business defect unless the task explicitly required that print output."
                    record["render_views"] = [{"view_id": original_view, "kind": "ORIGINAL_PRINT_DIAGNOSTIC_ONLY", "scope": "Unmodified default print layout; diagnostic only"}, {"view_id": wide_view, **proof}]
                    pages += wide_pages
                if extension in {".xlsx", ".xlsm"} and any(c["data_type"] == "f" for s in record.get("extraction", {}).get("sheets", []) for c in s["cells"]):
                    try:
                        record["derived_formula_recalculation"] = _recalculate_excel(copied, directory / "renders" / eid)
                    except Exception as exc:
                        record["derived_formula_recalculation_failure"] = type(exc).__name__
                        fail(eid + ":recalc-failure", record["locator"], exc, "derived-calculation")
            elif extension in {".html", ".htm"}:
                pages, runtime = _render_html(copied, directory / "renders" / eid, bundle)
                observations.append({"evidence_id": eid + ":browser", "artifact": record["filename"],
                                     "locator": record["locator"] + "/browser", "kind": "REAL_LOCAL_BROWSER_OBSERVATIONS",
                                     "text": canonical(runtime), "observations": runtime,
                                     "scope": "Fixed PC/tablet actual before/after interaction states; external network blocked. Coverage and every missing state are explicit; no backend services simulated"})
                record["evidence_ids"].append(eid + ":browser")
            elif extension in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
                pages = [{"path": str(copied.resolve()), "sha256": sha(copied), "locator": "image:1", "renderer": "Original-image-pixels"}]
            for n, page in enumerate(pages, 1):
                vid = f"{eid}:V{n:03}"
                visual.append({**page, "evidence_id": vid, "role": "output", "artifact": record["filename"], "relative_artifact": relative,
                               "locator": record["locator"] + "/" + page["locator"]})
                record["evidence_ids"].append(vid)
            if extension in {".html", ".htm"}:
                by_filename = {p.get("filename"): {"evidence_id": p["evidence_id"], "sha256": p["sha256"]} for p in visual if p.get("artifact") == record["filename"]}
                for state in runtime.get("state_catalog", []):
                    state["image_evidence_ids"] = list(dict.fromkeys(by_filename[x["filename"]]["evidence_id"] for x in state.get("image_refs", []) if x["filename"] in by_filename))
                    state["image_sha256"] = list(dict.fromkeys(x["sha256"] for x in state.get("image_refs", [])))
                observations[-1]["text"] = canonical(runtime)
            record["rendering_status"] = "COMPLETE" if pages else "NOT_APPLICABLE"
        except Exception as exc:
            record["rendering_status"] = "FAILED"
            fail(eid + ":render-failure", record["locator"], exc, "renderer")

    result = {"schema_version": SCHEMA, "validity": "EXPLORATORY_UNCALIBRATED", "task_id": task_rubric["task_id"],
              "blind_case_id": inventory_row["blind_case_id"], "binding": binding,
              "sources": source_records, "artifacts": artifacts, "outputs": artifacts,
              "visual_pages": visual, "render_images": visual, "functional_observations": observations,
              "missing_deliverables": missing, "evidence_failures": failures, "item_availability": {},
              "original_files_unchanged": all(p.is_file() and sha(p) == expected for p, expected in original_pins),
              "native_formula_policy": "Original formula/cache shown separately; never assume zero/missing/stale cached value is a wrong Actor calculation",
              "no_silent_truncation": True, "full_annotated_evidence_retained": True,
              "view_semantics_contract": "Original XLSX default print split is diagnosis only unless explicitly required. Sheet-wide view changes layout only, never data/style/width. Real HTML state evidence does not certify untested transitions. PDF original embedded fonts and PPTX declarative script slots are separate; platform compatibility remains unknown."}
    by_name = {a["filename"]: a for a in artifacts}
    # A uniquely corresponding wrong-named artifact remains content evidence.
    # The original filename and mismatch fact are retained for explicit naming criteria.
    missing_by_suffix, extras_by_suffix = {}, {}
    for a in artifacts:
        suffix = Path(a["filename"]).suffix.lower()
        if a["required"] and a["presence"] != "PRESENT":
            missing_by_suffix.setdefault(suffix, []).append(a)
        if not a["required"] and a["presence"] == "PRESENT":
            extras_by_suffix.setdefault(suffix, []).append(a)
    result["artifact_correspondence"] = []
    for suffix, absent in missing_by_suffix.items():
        candidates = extras_by_suffix.get(suffix, [])
        if len(absent) == 1 and len(candidates) == 1:
            wanted, actual = absent[0], candidates[0]
            by_name[wanted["filename"]] = actual
            result["artifact_correspondence"].append({"evidence_id": wanted["evidence_id"] + ":alias",
                "locator": wanted["locator"] + "/unique-format-correspondence", "expected_filename": wanted["filename"],
                "actual_filename": actual["filename"], "actual_evidence_id": actual["evidence_id"],
                "status": "UNIQUE_FORMAT_CONTENT_CANDIDATE_WITH_EXPLICIT_NAMING_MISMATCH",
                "text": "One unmatched required file and one delivered extra of the same native format; evaluate actual content, retain filename noncompliance separately. No TSR change."})
        elif candidates:
            for wanted in absent:
                result["artifact_correspondence"].append({"evidence_id": wanted["evidence_id"] + ":mapping-uncertain",
                    "locator": wanted["locator"] + "/uncertain-correspondence", "expected_filename": wanted["filename"],
                    "candidate_evidence_ids": [x["evidence_id"] for x in candidates], "status": "MAPPING_UNCERTAIN",
                    "text": "Multiple unmatched files cannot be paired by format alone; dependent quality content requires review rather than automatic zero."})
    modality_path = Path(task_rubric.get("modality_contract_path") or (RUNTIME / "item_modality_contract.json"))
    modalities = json.loads(modality_path.read_text(encoding="utf8")) if modality_path.is_file() else {}
    modalities = modalities.get("items", modalities)
    native_function_kinds = {"CALENDAR_NATIVE_PARSE", "ARCHIVE_LOCATOR_RESOLUTION", "MIME_NATIVE_PARSE_AND_HANDOFF", "MIME_AND_CALENDAR_PARSE",
        "MIME_ATTACHMENT_DECODE", "MIME_ZIP_PARSE_AND_HANDOFF", "PDF_TEXT_AND_ANNOTATION_INSPECTION", "SENSITIVE_CONTENT_NATIVE_SCAN",
        "PDF_GEOMETRY_AND_TEXT_INSPECTION", "PDF_READING_SEARCH_INSPECTION", "CALENDAR_AND_CSV_TIME_COMPARISON", "CALENDAR_RECURRENCE_EXPANSION"}
    for record in artifacts:
        if record.get("parsing_status") == "COMPLETE":
            relevant_items = [x["item_id"] for x in task_rubric["items"] if record["filename"] in x.get("deliverables", [])
                              and modalities.get(x["item_id"], {}).get("functional_evidence_kind") in native_function_kinds]
            if relevant_items:
                eid = record["evidence_id"] + ":native-inspection"
                observations.append({"evidence_id": eid, "artifact": record["filename"], "locator": record["locator"] + "/native-inspection",
                    "kind": "ACTUAL_TRUSTED_NATIVE_PARSER_INSPECTION", "item_ids": relevant_items,
                    "text": "The trusted local parser successfully read the complete selected native bytes. Actual structures/text/metadata are at " + record["evidence_id"] + ". No remote service, sending, or application import is asserted.",
                    "actual_native_sha256": record["sha256"], "parser": record.get("extraction", {}).get("parser")})
                record["evidence_ids"].append(eid)
    source_failure = any(s["parsing_status"] in {"FAILED", "INCOMPLETE"} for s in source_records)
    for item in task_rubric["items"]:
        dependencies = [by_name.get(x) for x in item.get("deliverables", [])]
        dependencies = [x for x in dependencies if x]
        available = [x for x in dependencies if x["presence"] == "PRESENT"]
        refs = [x["evidence_id"] for x in dependencies]
        status, reason = "AVAILABLE", "Full source and relevant native output evidence supplied; score still requires actual item-specific evidence"
        required_modes = modalities.get(item["item_id"], {}).get("required_evidence_modes", [x for x in item.get("evidence_modes", []) if x in {"visual", "functional"}])
        if dependencies and not available:
            status, reason = "ABSENT_REQUIRED_CONTENT", "All listed dependent deliverables are absent or zero bytes; fixed weights retained"
        elif source_failure:
            status, reason = "UNSCORED", "At least one source failed complete parsing; unknown evidence must not become zero"
        elif any(x.get("parsing_status") != "COMPLETE" for x in available):
            status, reason = "UNSCORED", "A dependent native artifact could not be completely parsed"
        elif "visual" in required_modes:
            need = [x for x in available if Path(x["filename"]).suffix.lower() in {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".html", ".htm", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".svg"}]
            if not need or any(x.get("rendering_status") != "COMPLETE" for x in need):
                status, reason = "UNSCORED", "Required native visual pages unavailable; text/metadata cannot substitute for layout pixels"
            else:
                refs += [v["evidence_id"] for v in visual if v["artifact"] in {x["filename"] for x in need}]
        if status == "AVAILABLE" and "functional" in required_modes:
            relevant = [o for o in observations if o.get("artifact") in {x["filename"] for x in available}]
            functional_kind = modalities.get(item["item_id"], {}).get("functional_evidence_kind")
            untested_specific_kinds = {"LOCAL_BROWSER_QUOTE_FLOW", "LOCAL_BROWSER_FILTER_AND_SORT", "LOCAL_BROWSER_QUOTE_SAMPLE_CALCULATION",
                                       "LOCAL_BROWSER_PRACTICE_REQUESTS", "CALENDAR_LOCAL_IMPORT_VALIDATION", "SEARCH_INSTRUCTION_REPLAY"}
            if functional_kind in untested_specific_kinds:
                relevant = []
            if relevant:
                refs += [o["evidence_id"] for o in relevant]
            else:
                status, reason = "UNSCORED", "No actual functional test evidence yet; source presence/self-reported success is insufficient"
        if any(c.get("status") == "MAPPING_UNCERTAIN" and c["expected_filename"] in item.get("deliverables", []) for c in result["artifact_correspondence"]):
            status, reason = "UNSCORED", "Unmatched same-format artifact correspondence is ambiguous; no filename-based content zero"
        result["item_availability"][item["item_id"]] = {"status": status, "reason": reason, "evidence_ids": list(dict.fromkeys(refs)),
                                                       "required_evidence_modes": required_modes,
                                                       "authored_binding_v3": item.get("evidence_binding_v3"),
                                                       "functional_evidence_kind": modalities.get(item["item_id"], {}).get("functional_evidence_kind"),
                                                       "sufficient_evidence_modes": [x for x in item.get("evidence_modes", []) if x not in {"visual", "functional"}]}
    try:
        from evidence_process import augment_process_evidence
    except ImportError:
        augment_process_evidence = None
    if augment_process_evidence:
        augmented = augment_process_evidence(task_rubric, inventory_row, result)
        if augmented is not None:
            result = augmented
        historical = result.get("historical_process_evidence", {})
        native = historical.get("native_trace", {})
        public_actions = native.get("public_action_evidence", [])
        for action in public_actions:
            action.setdefault("locator", "retained-session:line:" + str(action.get("source_line", "unknown")))
            action.setdefault("kind", "HISTORICAL_PUBLIC_TOOL_CALL_AND_FEEDBACK")
            action.setdefault("text", canonical({k: v for k, v in action.items() if k != "text"}))
            result["functional_observations"].append(action)
        protected = historical.get("protected_inputs", [])
        protection_eid = "PROCESS:FINAL_INPUT_SNAPSHOTS"
        if protected:
            result["functional_observations"].append({"evidence_id": protection_eid, "locator": "retained-attempt:final-input-byte-snapshots",
                "kind": "ACTUAL_FINAL_PROTECTED_INPUT_SHA_OBSERVATION", "protected_inputs": protected,
                "text": canonical(protected), "scope": "Final exported bytes only; a matching final SHA does not establish uninterrupted immutability"})
        for item in task_rubric["items"]:
            current = result["item_availability"][item["item_id"]]
            if "functional" in current.get("required_evidence_modes", []) and public_actions:
                current["evidence_ids"] += [x["evidence_id"] for x in public_actions]
                if current.get("functional_evidence_kind") == "PROTECTED_SOURCE_OBSERVATION" and protected and all(x["status"] in {"FINAL_BYTES_MATCH_SOURCE", "FINAL_BYTES_DIFFER_FROM_SOURCE"} for x in protected) and current["status"] == "UNSCORED" and current["reason"].startswith("No actual functional"):
                    current["status"] = "AVAILABLE"
                    current["evidence_ids"].append(protection_eid)
                    current["reason"] = "Actual final input SHA comparison and retained real tool feedback supplied; final equality is not uninterrupted immutability"
            process = historical.get("process_item_availability", {}).get(item["item_id"])
            if process and not process.get("history_available") and current["status"] == "AVAILABLE":
                current["status"] = "UNSCORED"
                current["reason"] = "Required real operation history is unavailable; artifact self-report is insufficient"
    if not result["original_files_unchanged"]:
        raise ValueError("original_file_changed_during_evidence_preparation")
    result["evidence_atoms"] = _evidence_atoms(result)
    result["atom_contract"] = {"representation": "JSON_POINTER_ONLY_NO_SECOND_TEXT_COPY", "resolves_against": "THIS_COMPLETE_EVIDENCE_JSON",
                               "string_quote": "literal scalar string", "other_scalar_quote": "JSON scalar representation",
                               "input_output_observation_roles_are_separate": True,
                               "no_atom_grants_semantic_support_by_itself": True}
    write(target, result)
    return result
