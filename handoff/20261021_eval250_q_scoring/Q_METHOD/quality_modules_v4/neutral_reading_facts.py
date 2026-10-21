"""Private, read-only cash-flow table reading facts; no verdict or score.

API: build_reading_facts(actual_xlsx_path, case) -> JSON-serializable dict.
All amounts are recomputed from the supplied SHA-bound inputs at call time.
This supports recognized row/column labels; other layouts return UNKNOWN rather
than assuming a wrong result. No static expected numbers or audit JSON is read.
"""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import ast,collections,hashlib,re
import openpyxl

def _sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def _d(v):return Decimal(str(v or 0))
def _money(v):return str(_d(v).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))

def _value(sheet,address,trail=()):
 if address in trail:raise ValueError('CIRCULAR_FORMULA')
 value=sheet[address].value
 if isinstance(value,(int,float,Decimal)):return _d(value)
 if not isinstance(value,str) or not value.startswith('='):raise ValueError('NOT_NUMERIC_OR_SUPPORTED_FORMULA')
 node=ast.parse(value[1:],mode='eval')
 def visit(n):
  if isinstance(n,ast.Expression):return visit(n.body)
  if isinstance(n,ast.Name) and re.fullmatch('[A-Z]+[0-9]+',n.id):return _value(sheet,n.id,trail+(address,))
  if isinstance(n,ast.Constant) and isinstance(n.value,(int,float)):return _d(n.value)
  if isinstance(n,ast.BinOp) and isinstance(n.op,(ast.Add,ast.Sub)):
   return visit(n.left)+visit(n.right) if isinstance(n.op,ast.Add) else visit(n.left)-visit(n.right)
  if isinstance(n,ast.UnaryOp) and isinstance(n.op,ast.USub):return -visit(n.operand)
  raise ValueError('UNSUPPORTED_FORMULA')
 return visit(node)

def build_reading_facts(actual_xlsx_path,case):
 path=Path(actual_xlsx_path);original_sha=_sha(path)
 facts={'schema':'cashflow-private-neutral-reading-facts-v1','artifact_sha256':original_sha,'scope':'PRIVATE_JUDGE_READBACK_NOT_ACTOR_INPUT','availability':'UNKNOWN','column_map':[],'cells':[],'source_records':[],'source_inputs':[],'limitations':['Only unambiguous recognized original table labels are mapped.','Signed expenses remain signed; negative values are not labelled wrong.','Numeric/relational facts do not assign rubric scores or business pass/fail.','Unsupported layouts or formulas stay UNKNOWN; no unique layout requirement.']}
 try:
  rules=[];source_rows=[]
  for item in case['inputs']:
   p=Path(item['path']);actual=_sha(p)
   if actual!=item['sha256']:raise ValueError('SOURCE_SHA_MISMATCH')
   facts['source_inputs'].append({'name':item['name'],'sha256':actual,'source_id':item.get('source_id')})
   if p.suffix.lower()=='.md':rules.append(p.read_text('utf-8-sig'))
   if p.suffix.lower()!='.xlsx':continue
   wb=openpyxl.load_workbook(p,read_only=True,data_only=False)
   try:
    sheet=wb.active
    if [c.value for c in sheet[1]]!=['流水ID','日期','摘要','收入元','支出元','对方机构']:continue
    company=p.stem.removesuffix('流水')
    for row in list(sheet.iter_rows())[1:]:
     identifier,date,summary,income,expense,counterparty=[c.value for c in row]
     if not identifier:continue
     if _d(income)!=0 and _d(expense)!=0:raise ValueError('AMBIGUOUS_TWO_SIDED_SOURCE_ROW')
     if not re.match(r'\d{4}-\d{2}-\d{2}',str(date)):raise ValueError('UNSUPPORTED_SOURCE_DATE')
     source_rows.append({'source_ref':str(item.get('source_id',p.name))+':row:'+str(row[0].row),'source_name':item['name'],'source_sheet':sheet.title,'source_row':row[0].row,'company':company,'id':str(identifier),'date':str(date),'period':str(date)[:7],'summary':str(summary),'income':_money(income),'expense':_money(expense),'signed':_money(_d(income)-_d(expense))})
   finally:wb.close()
  rule='\n'.join(rules)
  if not source_rows:raise ValueError('NO_RECOGNIZED_SOURCE_LEDGERS')
  known=['产品销售回款','研发补贴','银行利息收入','内部往来收款','材料采购','员工报销','备用金转出','服务外包','内部往来付款','设备采购','项目投资款','股东增资']
  if any(x['summary'] not in known or x['summary'] not in rule for x in source_rows):raise ValueError('UNSUPPORTED_RULE_OR_SOURCE_CLASSIFICATION')
  components={x:{x} for x in known}
  components.update({'利息收入':{'银行利息收入'},'经营流入小计':set(known[:4]),'经营流出小计':set(known[4:9]),'投资流出小计':set(known[9:11]),'筹资流入小计':{'股东增资'},'经营净额':set(known[:9]),'投资净额':set(known[9:11]),'筹资净额':{'股东增资'},'总净额':set(known)})
  # Only the explicitly named matching pair(s) in the actual original rule may
  # contribute elimination. Equal amounts or repeated IDs do not establish it.
  pairing=re.compile(r'内部匹配编号([^：\s]+)：(.+?)([A-Za-z]+\d+)收入([\d.]+)元对应(.+?)([A-Za-z]+\d+)支出([\d.]+)元')
  offset_rows=[]
  for match in pairing.finditer(rule):
   number,inc_company,inc_id,inc_amount,out_company,out_id,out_amount=match.groups()
   pairs=[x for x in source_rows if (x['company'],x['id']) in [(inc_company,inc_id),(out_company,out_id)]]
   if len(pairs)!=2 or _d(inc_amount)!=_d(out_amount):raise ValueError('AMBIGUOUS_EXPLICIT_INTERNAL_PAIR')
   for source in pairs:
    expected=_d(inc_amount) if source['company']==inc_company else -_d(out_amount)
    if _d(source['signed'])!=expected:raise ValueError('EXPLICIT_PAIR_SOURCE_AMOUNT_MISMATCH')
    offset_rows.append({**source,'pair_id':number,'source_signed':source['signed'],'signed':_money(-_d(source['signed'])),'source_multiplier':-1})
  if '内部匹配编号' in rule and not offset_rows:raise ValueError('UNSUPPORTED_INTERNAL_PAIR_RULE')
  facts['source_records']=source_rows;facts['explicit_elimination_terms']=offset_rows
  companies=sorted({x['company'] for x in source_rows});periods=sorted({x['period'] for x in source_rows})
  if len({p[:4] for p in periods})!=1:raise ValueError('AMBIGUOUS_MULTIYEAR_MONTH_HEADER')
  year=periods[0][:4];wb=openpyxl.load_workbook(path,data_only=False)
  try:
   for sheet in wb:
    # The header is observed, not assigned to a preferred coordinate. A source
    # entity/month table requires a project/item label in its own first column.
    candidates=[row for row in range(1,min(sheet.max_row,12)+1) if str(sheet.cell(row,1).value or '').strip()=='项目']
    if len(candidates)!=1:continue
    header=candidates[0];mapping={}
    for column in range(2,sheet.max_column+1):
     value=str(sheet.cell(header,column).value or '').strip()
     if value in companies:entity=value;period=None
     elif value=='集团合计':entity='集团';period=None
     else:
      found=re.fullmatch('('+'|'.join(re.escape(x) for x in companies+['内部抵销','集团'])+r')(\d{1,2})月',value)
      if not found:continue
      entity,month=found.groups();period=year+'-'+month.zfill(2)
      if period not in periods:continue
     mapping[column]={'sheet':sheet.title,'column':openpyxl.utils.get_column_letter(column),'header_cell':sheet.cell(header,column).coordinate,'header':value,'entity':entity,'period':period}
    facts['column_map'].extend(mapping.values())
    for row in range(header+1,sheet.max_row+1):
     label=str(sheet.cell(row,1).value or '').strip();keys=components.get(label)
     if not keys:continue
     for column,column_info in mapping.items():
      cell=sheet.cell(row,column)
      if cell.value is None:continue
      entity,period=column_info['entity'],column_info['period']
      relevant=[r for r in source_rows if r['summary'] in keys and (period is None or r['period']==period) and (entity=='集团' or r['company']==entity)] if entity!='内部抵销' else []
      offsets=[r for r in offset_rows if r['summary'] in keys and (period is None or r['period']==period)] if entity in ('集团','内部抵销') else []
      grouped=collections.defaultdict(list)
      for r in relevant:grouped[(r['company'],r['period'],r['summary'],'SOURCE')].append(r)
      for r in offsets:grouped[(r['company'],r['period'],r['summary'],'EXPLICIT_ELIMINATION')].append(r)
      terms=[]
      for (company,source_period,category,kind),records in sorted(grouped.items()):
       terms.append({'kind':kind,'company':company,'period':source_period,'source_summary':category,'signed_value':_money(sum((_d(r['signed']) for r in records),Decimal(0))),'source_record_refs':[r['source_ref'] for r in records],'multiplier':-1 if kind=='EXPLICIT_ELIMINATION' else 1})
      independent=sum((_d(t['signed_value']) for t in terms),Decimal(0))
      original_formula=cell.value if cell.data_type=='f' else None
      dependencies=[]
      if original_formula:
       for addr in sorted(set(re.findall(r'(?<![A-Za-z_])\$?([A-Z]+)\$?(\d+)',original_formula))):
        address=''.join(addr);target=sheet[address]
        dependencies.append({'sheet':sheet.title,'cell':address,'raw':target.value,'row_label':str(sheet.cell(target.row,1).value or '').strip(),'column_label':str(sheet.cell(header,target.column).value or '').strip()})
      try:value=_money(_value(sheet,cell.coordinate));formula_support='SUPPORTED_SAME_SHEET_ADD_SUBTRACT'
      except (ValueError,SyntaxError,TypeError):value=None;formula_support='UNKNOWN_UNSUPPORTED_FORMULA'
      relation=[]
      if entity=='集团':
       desired=[(x,period) for x in companies+['内部抵销']] if period is not None else [('集团',p) for p in periods]
       related=[sheet.cell(row,c) for entity_period in desired for c,m in mapping.items() if (m['entity'],m['period'])==entity_period and c!=column]
       if len(related)==len(desired):
        relation=[{'sheet':sheet.title,'cell':t.coordinate,'raw':t.value,'column_label':str(sheet.cell(header,t.column).value or '').strip()} for t in related]
      facts['cells'].append({'sheet':sheet.title,'cell':cell.coordinate,'row_label':label,'column_label':column_info['header'],'raw':cell.value,'formula':original_formula,'value':value,'formula_value_method':formula_support,'formula_dependency_cells':dependencies,'same_row_observed_components':relation,'source_terms':terms,'independent_source_value':_money(independent),'source_relation':'Sum of source_terms signed_value, preserving original receipt/payment convention and only explicitly paired eliminations.'})
  finally:wb.close()
  if not facts['cells']:raise ValueError('NO_UNAMBIGUOUS_SUMMARY_TABLES')
  if _sha(path)!=original_sha:raise ValueError('ARTIFACT_CHANGED_DURING_READ')
  for item in case['inputs']:
   if _sha(item['path'])!=item['sha256']:raise ValueError('SOURCE_CHANGED_DURING_READ')
  facts['availability']='AVAILABLE';return facts
 except Exception as exc:
  facts['availability']='UNKNOWN';facts['unavailable_reason']=type(exc).__name__+': '+str(exc)
  return facts
