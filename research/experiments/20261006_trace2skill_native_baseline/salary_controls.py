"""Predetermined synthetic new tasks; not extra historical traces or benchmarks."""
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from audited_runtime import dump,sha,bootstrap_imports

def prepare(root):
    bootstrap_imports()
    import openpyxl
    from analysis.evaluate_output import compare_workbooks
    root=Path(root); root.mkdir()
    cases={'salary-new-1':[123456,98765], 'salary-new-2':[246913,158001]}
    rows=[]; checks=[]
    for task_id,annuals in cases.items():
        folder=root/task_id; folder.mkdir()
        wb=openpyxl.Workbook(); ws=wb.active; ws.title='SALARY'
        ws.append(['Employee','Annual yuan','Monthly yuan'])
        for idx,annual in enumerate(annuals,1): ws.append([f'Employee {idx}',annual,None])
        wb.save(folder/f'1_{task_id}_input.xlsx')
        expected=[str((Decimal(a)/Decimal(12)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)) for a in annuals]
        for idx,value in enumerate(expected,2): ws.cell(idx,3,float(value))
        gold=folder/f'1_{task_id}_golden.xlsx'; wb.save(gold)
        positive=folder/'control_positive.xlsx'; wb.save(positive)
        ws['C2']=0; negative=folder/'control_negative.xlsx'; wb.save(negative); wb.close()
        answer='SALARY!C2:C3'
        pos=compare_workbooks(str(gold),str(positive),answer)
        neg=compare_workbooks(str(gold),str(negative),answer)
        if not pos[0] or neg[0]: raise RuntimeError('Synthetic salary scorer controls failed')
        # Control files are gold-facing only; native input finder selects *_input.xlsx.
        rows.append({'id':task_id,'spreadsheet_path':task_id,'instruction_type':'Cell-Level Manipulation',
            'answer_position':answer,'answer_sheet':'SALARY',
            'instruction':'请把SALARY表每名员工的年工资（B列，单位元）按一年12个月折算为月工资，'
                          '填入C2:C3，保留两位小数。仅处理此换算，不推断人数、其他费用或现金流。保留其余数据。'})
        checks.append({'id':task_id,'annual_yuan':annuals,'expected_monthly_yuan':expected,
            'positive_native_score':pos,'negative_native_score':neg,
            'input_sha256':sha(folder/f'1_{task_id}_input.xlsx'),'gold_sha256':sha(gold)})
    dump(root/'dataset.json',rows)
    dump(root/'protocol.json',{'purpose':'Functional reuse acceptance only, NOT statistical skill benefit',
        'origin':'Researcher-constructed independent new values, frozen before any task model call',
        'no_historical_task_reexecution':True,'no_learning_from_acceptance_tasks':True,
        'scorer':'Unmodified author analysis.evaluate_output.compare_workbooks','checks':checks})
    return rows
