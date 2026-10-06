"""Free, local positive/negative controls for the unchanged Trace2Skill runtime.

This script never creates a model client, executes a business rollout, installs a
dependency, reads credentials, or changes the baseline checkout. All control
outputs stay under this experiment's private/runtime_preflight directory.
"""
from __future__ import annotations

import contextlib
import datetime
import hashlib
import importlib
import importlib.metadata
import importlib.util
import io
import json
import os
from pathlib import Path
import platform
import shutil
import subprocess
import sys
import traceback

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
PROJECT = HERE.parents[2]
BASELINE = PROJECT / "research/baselines/Trace2Skill"
CONTROL_ROOT = HERE / "private/runtime_preflight"
RUNTIME = CONTROL_ROOT / "runtime"
CHECKER_SOURCE = PROJECT / "enginering/demo/.runtime/node_modules/openclaw/skills/skill-creator/scripts/quick_validate.py"
DATASET = BASELINE / "data/spreadsheetbench_verified/spreadsheetbench_verified_400"
REPORT = HERE / "runtime_readiness.json"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def decode(value: bytes | str | None) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if b"\x00" in value[:100]:
        return value.decode("utf-16le", errors="replace").lstrip("\ufeff")
    return value.decode("utf-8", errors="replace")


def command(args: list[str], timeout: int = 12) -> dict:
    """No shell, no secret-bearing arguments, and no environment dump."""
    try:
        result = subprocess.run(args, cwd=RUNTIME, capture_output=True, timeout=timeout)
        return {"args": args, "returncode": result.returncode,
                "stdout": decode(result.stdout)[-6000:], "stderr": decode(result.stderr)[-3000:]}
    except subprocess.TimeoutExpired as exc:
        return {"args": args, "status": "TIMEOUT", "stdout": decode(exc.stdout)[-2000:],
                "stderr": decode(exc.stderr)[-2000:]}
    except (OSError, ValueError) as exc:
        return {"args": args, "status": "UNAVAILABLE", "error": str(exc)}


@contextlib.contextmanager
def cwd(path: Path):
    before = Path.cwd()
    os.chdir(path)
    try:
        yield
    finally:
        os.chdir(before)


def capture(label: str, function):
    stdout, stderr = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
        result = function()
    (CONTROL_ROOT / f"{label}.stdout.txt").write_text(stdout.getvalue(), encoding="utf-8")
    (CONTROL_ROOT / f"{label}.stderr.txt").write_text(stderr.getvalue(), encoding="utf-8")
    return result


def dependency_inventory() -> dict:
    packages = {}
    for name in ("openai", "requests", "tqdm", "openpyxl", "diskcache", "yaml", "evaluation_official"):
        try:
            spec = importlib.util.find_spec(name)
            item = {"available": spec is not None, "origin": spec.origin if spec else None}
            try:
                item["version"] = importlib.metadata.version("PyYAML" if name == "yaml" else name)
            except importlib.metadata.PackageNotFoundError:
                item["version"] = None
            packages[name] = item
        except Exception as exc:
            packages[name] = {"available": False, "error": str(exc)}
    return {"python": sys.executable, "version": platform.python_version(), "packages": packages}


def environment_inventory() -> dict:
    tools = {name: shutil.which(name) for name in
             ("bash", "python", "python3", "soffice", "libreoffice", "timeout", "wsl", "docker")}
    probes = {}
    git_bash = Path("C:/Program Files/Git/bin/bash.exe")
    if git_bash.is_file():
        probes["git_bash_version"] = command([str(git_bash), "--version"])
        probes["git_bash_programs"] = command([str(git_bash), "-c", "command -v python python3 timeout soffice"])
        probes["git_bash_python_identity"] = command([
            str(git_bash), "-c", "python -B -c 'import sys; print(sys.executable)'"
        ])
    if tools["wsl"]:
        probes["wsl_distributions"] = command([tools["wsl"], "--list", "--verbose"])
    if tools["docker"]:
        probes["docker_daemon_version"] = command([tools["docker"], "info", "--format", "{{.ServerVersion}}"])
    if tools["soffice"]:
        probes["soffice_version"] = command([tools["soffice"], "--version"])
    # Exercise the author's actual execution tool, without a model or business command.
    from spreadsheet_agent.tools.bash import create_bash_tool
    bash_output = create_bash_tool(str(RUNTIME), timeout=5).execute(command="command -v python")
    return {"platform": platform.platform(), "native_default_shell": "cmd.exe" if os.name == "nt" else "/bin/sh",
            "paths": tools, "probes": probes, "author_bash_probe": {
                "command": "command -v python", "result": bash_output,
                "note": "Author shell=True does not select Git Bash. This is an environment probe, not a failed task."
            }, "linux_runtime_certified": False,
            "libreoffice_recalculation_certified": False}


def checker_controls() -> dict:
    target = RUNTIME / "skills/skill-creator/scripts/quick_validate.py"
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(CHECKER_SOURCE, target)
    from skill_evolver.skill_evolving_agent import SkillEvolver, QUICK_VALIDATE_SCRIPT
    results = {}
    texts = {
        "valid": "---\nname: runtime-format-control\ndescription: Local format control only.\n---\n# Format control\n\nThis is not a generated research skill.\n",
        "invalid": "---\nname: runtime-format-control\n---\n# Missing description control\n"
    }
    for label, text in texts.items():
        skill = CONTROL_ROOT / "checker" / label
        skill.mkdir(parents=True, exist_ok=True)
        (skill / "SKILL.md").write_text(text, encoding="utf-8")
        # validate_skill only needs skill_dir. Bypass client-bearing initialization.
        evolver = SkillEvolver.__new__(SkillEvolver)
        evolver.skill_dir = skill
        with cwd(RUNTIME):
            if not QUICK_VALIDATE_SCRIPT.is_file():
                raise RuntimeError("Configured original relative checker still missing")
            outcome = evolver.validate_skill()
        direct = command([sys.executable, "-B", str(target), str(skill)])
        results[label] = {"author_validate_skill": {"passed": outcome[0], "message": outcome[1]},
                          "checker_cli": direct, "skill_file_sha256": sha(skill / "SKILL.md")}
    passed = (results["valid"]["author_validate_skill"]["passed"] is True
              and results["invalid"]["author_validate_skill"]["passed"] is False
              and results["valid"]["checker_cli"].get("returncode") == 0
              and results["invalid"]["checker_cli"].get("returncode") == 1)
    return {"status": "PASS" if passed else "FAIL", "source": str(CHECKER_SOURCE), "source_sha256": sha(CHECKER_SOURCE),
            "configured_cwd": str(RUNTIME), "relative_path": str(QUICK_VALIDATE_SCRIPT),
            "copied_path": str(target), "copy_sha256": sha(target), "controls": results,
            "scope": "Format only; no Markdown consistency, semantic quality or executable skill certification."}


def scoring_controls() -> dict:
    import openpyxl
    from analysis import evaluate_output
    from analysis.error_analysis_agent import create_evaluate_tool
    import evaluate_with_official
    rows = json.loads((DATASET / "dataset.json").read_text(encoding="utf-8"))
    index = next(i for i, row in enumerate(rows) if str(row["id"]) == "13-1")
    row = rows[index]
    source = DATASET / row["spreadsheet_path"]
    gold = source / "1_13-1_golden.xlsx"
    initial = source / "1_13-1_init.xlsx"
    wb = openpyxl.load_workbook(gold, data_only=True)
    original_scope = row["answer_position"]
    scope = original_scope
    # Explicit metadata-only adapter: the author's bare-range fallback selects
    # the first worksheet and otherwise silently ignores answer_sheet.
    if "!" not in scope and row.get("answer_sheet"):
        scope = "'" + row["answer_sheet"] + "'!" + scope
    first_scope = scope.split(",")[0].strip()
    if "!" in first_scope:
        sheet, cell_range = first_scope.split("!")
        sheet = sheet.strip("'")
    else:
        sheet, cell_range = wb.sheetnames[0], first_scope
    cell = evaluate_output.generate_cell_names(cell_range.strip("'"))[0]
    original_value = wb[sheet][cell].value
    wb.close()
    scoring_dataset = CONTROL_ROOT / "scorer_metadata_adapter"
    copied_source = scoring_dataset / row["spreadsheet_path"]
    copied_source.mkdir(parents=True, exist_ok=True)
    for source_file in (gold, initial):
        shutil.copyfile(source_file, copied_source / source_file.name)
    adapted_row = dict(row)
    adapted_row["answer_position"] = scope
    write_json(scoring_dataset / "dataset.json", [adapted_row])
    controls = {}
    for label in ("correct", "incorrect"):
        root = CONTROL_ROOT / "scorer" / label
        output_dir = root / "outputs"
        out = output_dir / row["spreadsheet_path"] / "1_13-1_output.xlsx"
        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(gold, out)
        if label == "incorrect":
            mutated = openpyxl.load_workbook(out)
            mutated[sheet][cell] = "__DELIBERATE_CONTROL_MISMATCH__"
            mutated.save(out)
            mutated.close()
        cli = command([sys.executable, "-B", str(BASELINE / "analysis/evaluate_output.py"),
                       "--output_file", str(out), "--ground_truth", str(gold), "--answer_position", scope])
        flag = root / "evaluate_passed.flag"
        if flag.exists():
            flag.unlink()  # reset this explicitly owned control flag only
        pass_state = {"passed": False}
        tool = create_evaluate_tool(str(root), pass_state, flag)
        observation = tool.execute(output_file=str(out), ground_truth=str(gold), answer_position=scope)
        (root / "agentic_evaluation_observation.txt").write_text(observation, encoding="utf-8")
        batch = capture(f"evaluate_with_official_{label}", lambda: evaluate_with_official.evaluate(
            str(scoring_dataset), str(output_dir), start_idx=0, end_idx=1, verbose=True))
        write_json(root / "evaluate_with_official_results.json", batch)
        native_result = batch["results"][0]
        controls[label] = {"output_file": str(out), "output_sha256": sha(out),
                           "evaluate_output_cli": cli, "agentic_pass_state": pass_state,
                           "agentic_pass_flag_exists": flag.exists(), "agentic_observation": observation,
                           "evaluate_with_official_result": native_result, "summary": batch["summary"]}
    # Replay the untouched bare-range interface against a defect on the actual
    # requested worksheet. Preserve its false PASS rather than hide the bug.
    incorrect_output = Path(controls["incorrect"]["output_file"])
    bare_cli = command([sys.executable, "-B", str(BASELINE / "analysis/evaluate_output.py"),
                        "--output_file", str(incorrect_output), "--ground_truth", str(gold),
                        "--answer_position", original_scope])
    bare_batch = capture("evaluate_with_official_original_bare_scope", lambda: evaluate_with_official.evaluate(
        str(DATASET), str(CONTROL_ROOT / "scorer/incorrect/outputs"), start_idx=index, end_idx=index + 1, verbose=True))
    write_json(CONTROL_ROOT / "original_bare_scope_false_positive.json", bare_batch)
    passed = (controls["correct"]["evaluate_output_cli"].get("returncode") == 0
              and controls["incorrect"]["evaluate_output_cli"].get("returncode") == 1
              and controls["correct"]["agentic_pass_flag_exists"] is True
              and controls["incorrect"]["agentic_pass_flag_exists"] is False
              and controls["correct"]["evaluate_with_official_result"]["success"] is True
              and controls["incorrect"]["evaluate_with_official_result"]["success"] is False)
    return {"status": "PASS" if passed else "FAIL", "task_id": row["id"], "dataset_index": index,
            "dataset": str(DATASET / "dataset.json"), "dataset_sha256": sha(DATASET / "dataset.json"),
            "dataset_rows": len(rows), "original_answer_position": original_scope,
            "original_answer_sheet": row.get("answer_sheet"), "qualified_answer_position": scope, "effective_sheet": sheet,
            "metadata_only_adapter": {"dataset": str(scoring_dataset / "dataset.json"),
                "change": "answer_position explicitly qualified with the existing answer_sheet; original files and task wording unchanged",
                "copied_input_sha256": sha(copied_source / initial.name), "copied_gold_sha256": sha(copied_source / gold.name)},
            "original_bare_scope_negative_control": {"evaluate_output_cli": bare_cli,
                "evaluate_with_official_result": bare_batch["results"][0],
                "false_positive_observed": bare_cli.get("returncode") == 0 and bare_batch["results"][0]["success"] is True,
                "note": "Wrong LISTS output is accepted because bare A3:D32 is checked on first sheet RANGES. This is measurement plumbing, not algorithm quality."},
            "source_input": str(initial), "input_sha256": sha(initial), "source_gold": str(gold), "gold_sha256": sha(gold),
            "official_external_module_loaded": evaluate_with_official.official_compare_workbooks is not None,
            "active_compare_function_module": (evaluate_with_official.official_compare_workbooks.__module__
                if evaluate_with_official.official_compare_workbooks is not None
                else evaluate_with_official.local_compare_workbooks.__module__),
            "deliberate_negative_control": {"cell": f"{sheet}!{cell}", "original_value_type": type(original_value).__name__,
                "replacement": "__DELIBERATE_CONTROL_MISMATCH__", "note": "Constructed diagnostic; not a model rollout or evidence of algorithm failure. Saving may clear other formula caches."},
            "controls": controls, "formula_recalculation_tested": False,
            "scope": "Scoring and Agentic PASS plumbing only; no agent diagnosis, repair rollout or learned-skill benefit."}


def main() -> int:
    CONTROL_ROOT.mkdir(parents=True, exist_ok=True)
    RUNTIME.mkdir(parents=True, exist_ok=True)
    if REPORT.exists():
        prior = CONTROL_ROOT / "previous_reports" / (sha(REPORT) + ".json")
        prior.parent.mkdir(parents=True, exist_ok=True)
        if not prior.exists():
            shutil.copyfile(REPORT, prior)
    # The author's evaluator spawns Python without -X utf8. Native Windows may
    # otherwise emit CP936 bytes which its UTF-8 parent cannot decode.
    os.environ["PYTHONUTF8"] = "1"
    os.environ["PYTHONIOENCODING"] = "utf-8"
    os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
    sys.path.insert(0, str(BASELINE))
    sys.path.insert(0, str(BASELINE / "src"))
    source_paths = [BASELINE / name for name in (
        "skill_evolver/skill_evolving_agent.py", "skill_evolver/parallel_evolving_agent.py",
        "analysis/error_analysis_agent.py", "analysis/evaluate_output.py", "analysis/report_parsing.py",
        "spreadsheet_agent/tools/bash.py", "spreadsheet_agent/skills/xlsx/recalc.py",
        "evaluate_with_official.py", "spreadsheetbench_support.py")]
    before = {str(path.relative_to(BASELINE)): sha(path) for path in source_paths}
    report = {"created_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
              "protocol": "native-runtime-free-controls-v1", "script_sha256": sha(Path(__file__)),
              "baseline": str(BASELINE), "runtime_cwd": str(RUNTIME),
              "pythonpath": [str(BASELINE), str(BASELINE / "src")],
              "required_runtime_environment": {"PYTHONUTF8": "1", "PYTHONIOENCODING": "utf-8", "PYTHONDONTWRITEBYTECODE": "1"},
              "model_calls": 0, "business_rollouts": 0, "system_installations": 0,
              "source_hashes_before": before, "dependencies": dependency_inventory()}
    for label, function in (("checker", checker_controls), ("scorer", scoring_controls), ("environment", environment_inventory)):
        try:
            report[label] = function()
        except Exception as exc:
            report[label] = {"status": "ERROR", "error": str(exc), "traceback": traceback.format_exc()}
        write_json(REPORT, report)
    after = {str(path.relative_to(BASELINE)): sha(path) for path in source_paths}
    report["source_hashes_after"] = after
    report["baseline_inspected_sources_unchanged"] = before == after
    report["free_controls_passed"] = (report["checker"].get("status") == "PASS"
                                        and report["scorer"].get("status") == "PASS" and before == after)
    report["full_native_pipeline_ready"] = False
    report["remaining_gates"] = [
        "No model/rollout/Agentic diagnosis or minimal-repair end-to-end certification in this free preflight.",
        "Native Windows shell is not automatically Git Bash/Linux; actual agent command environment must be frozen.",
        "LibreOffice formula recalculation not certified; correct/incorrect controls use existing cached values.",
        "Scorer external-module versus author-local fallback is explicitly reported; do not silently claim the external module.",
        "Original weak self-create initialization and enterprise-domain adapter remain separate work.",
        "Format checker is copied OpenClaw version, not an unpublished author-original checker; provenance and hash are explicit."
    ]
    write_json(REPORT, report)
    print(json.dumps({"report": str(REPORT), "free_controls_passed": report["free_controls_passed"],
                      "checker": report["checker"].get("status"), "scorer": report["scorer"].get("status"),
                      "baseline_sources_unchanged": before == after, "model_calls": 0}, ensure_ascii=False))
    return 0 if report["free_controls_passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
