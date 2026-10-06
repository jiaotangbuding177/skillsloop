"""Project-local LibreOffice extraction and original recalc.py free controls.

Authorized runtime preparation only: no model, system installation, credentials,
author source edit, or historical command execution. Fresh outputs are exclusive.
"""
from __future__ import annotations

import json
import subprocess
import sys
import time
import traceback
import urllib.request
from decimal import Decimal
from hashlib import sha256
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
PREFLIGHT = OUT / "private/runtime_preflight"
RESOURCES = PREFLIGHT / "安装资源"
RUNTIME = PREFLIGHT / "libreoffice_runtime"
PROFILE = PREFLIGHT / "lo-profile"
CONTROLS = PREFLIGHT / "libreoffice_controls"
REPORT = OUT / "libreoffice_readiness.json"
RECALC = ROOT / "research/baselines/Trace2Skill/spreadsheet_agent/skills/xlsx/recalc.py"
PACKAGE = "LibreOffice_26.8.0_Win_x86-64.msi"
URL = "https://download.documentfoundation.org/libreoffice/stable/26.8.0/win/x86_64/" + PACKAGE
SIZE = 374906880
HASH = "4aa6c6e1895f4055104effcb556bd3362d20c6ad707c149543304f395ef9db95"


def digest(path):
    value = sha256()
    with Path(path).open("rb") as stream:
        for data in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(data)
    return value.hexdigest()


def rel(path):
    return Path(path).relative_to(ROOT).as_posix()


def write(path, data):
    with Path(path).open("x", encoding="utf-8") as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def download():
    target = RESOURCES / PACKAGE
    if target.exists():
        assert target.stat().st_size == SIZE and digest(target) == HASH
        return target, {"status": "EXISTING_VERIFIED_PACKAGE_REUSED", "bytes": SIZE, "sha256": HASH}
    temporary = RESOURCES / (PACKAGE + ".part")
    if temporary.exists():
        raise RuntimeError("Previous partial download exists; preserve it rather than silently retry or overwrite")
    RESOURCES.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(URL, headers={"User-Agent": "Trace2Skill-local-runtime-preflight"})
    total, next_notice, h = 0, 32 * 1024 * 1024, sha256()
    print("Downloading official MSI; expected bytes=" + str(SIZE), flush=True)
    with urllib.request.urlopen(request, timeout=30) as response, temporary.open("xb") as stream:
        final_url = response.geturl()
        for block in iter(lambda: response.read(1024 * 1024), b""):
            stream.write(block)
            h.update(block)
            total += len(block)
            if total >= next_notice:
                print("Downloaded bytes=" + str(total), flush=True)
                next_notice += 32 * 1024 * 1024
    observed = h.hexdigest()
    if total != SIZE or observed != HASH:
        raise RuntimeError("Official MSI size/hash mismatch; partial package retained, never executed")
    assert not target.exists()
    temporary.rename(target)
    return target, {"status": "DOWNLOADED_SIZE_AND_OFFICIAL_SHA256_VERIFIED", "final_public_url": final_url,
                    "bytes": total, "sha256": observed, "path": rel(target)}


def run_control(kind, formula, expected, program):
    target = CONTROLS / (kind + ".xlsx")
    assert not target.exists()
    workbook = openpyxl.Workbook()
    workbook.active.title = "RecalcControl"
    workbook.active["A1"] = formula
    workbook.save(target)
    workbook.close()
    initial_hash = digest(target)
    before = openpyxl.load_workbook(target, data_only=True)
    before_cache = before.active["A1"].value
    before.close()
    # Minimal configuration environment, never copy/read API-key variables.
    package_site = str(Path(openpyxl.__file__).resolve().parent.parent)
    temp = PROFILE / "temporary"
    temp.mkdir(parents=True, exist_ok=True)
    env = {
        "SystemRoot": "C:\\WINDOWS", "WINDIR": "C:\\WINDOWS",
        "COMSPEC": "C:\\WINDOWS\\system32\\cmd.exe",
        "PATH": str(program) + ";C:\\Python314;C:\\WINDOWS\\system32;C:\\WINDOWS",
        "USERPROFILE": str(PROFILE), "APPDATA": str(PROFILE / ".config"),
        "LOCALAPPDATA": str(PROFILE / ".local"), "TEMP": str(temp), "TMP": str(temp),
        "PYTHONPATH": package_site, "PYTHONIOENCODING": "utf-8", "PYTHONUTF8": "1",
    }
    command = [sys.executable, "-B", "-X", "utf8", str(RECALC), str(target), "30"]
    print("Original recalc.py control: " + kind, flush=True)
    started = time.monotonic()
    process = subprocess.Popen(command, cwd=str(CONTROLS), env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               text=True, encoding="utf-8", errors="replace")
    timed_out = False
    try:
        stdout, stderr = process.communicate(timeout=90)
    except subprocess.TimeoutExpired:
        timed_out = True
        # Kill only this explicitly owned control's process tree.
        subprocess.run(["C:\\WINDOWS\\system32\\taskkill.exe", "/PID", str(process.pid), "/T", "/F"],
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=20)
        stdout, stderr = process.communicate(timeout=10)
    with (CONTROLS / (kind + ".stdout.txt")).open("x", encoding="utf-8") as stream:
        stream.write(stdout)
    with (CONTROLS / (kind + ".stderr.txt")).open("x", encoding="utf-8") as stream:
        stream.write(stderr)
    try:
        parsed = json.loads(stdout)
    except json.JSONDecodeError:
        parsed = None
    cached = openpyxl.load_workbook(target, data_only=True)
    cell = cached.active["A1"]
    actual, actual_type = cell.value, cell.data_type
    cached.close()
    formulas = openpyxl.load_workbook(target, data_only=False)
    formula_after = formulas.active["A1"].value
    formulas.close()
    if kind == "positive":
        correct_cache = isinstance(actual, (int, float)) and abs(Decimal(str(actual)) - expected) < Decimal("0.000001")
        correct_report = isinstance(parsed, dict) and parsed.get("status") == "success" and parsed.get("total_errors") == 0
    else:
        correct_cache = actual == "#DIV/0!" and actual_type == "e"
        correct_report = isinstance(parsed, dict) and parsed.get("status") == "errors_found" and parsed.get("total_errors") == 1 and "#DIV/0!" in parsed.get("error_summary", {})
    passed = process.returncode == 0 and not timed_out and correct_cache and correct_report and formula_after == formula
    return {"control": kind, "constructed_not_enterprise_episode": True, "file": rel(target),
            "formula": formula, "initial_cached_value": before_cache, "cached_value_after": actual,
            "cached_cell_type_after": actual_type, "formula_after": formula_after,
            "expected": str(expected), "original_recalc_report": parsed,
            "exit_code": process.returncode, "outer_timeout": timed_out,
            "outer_timeout_seconds": 90, "elapsed_seconds": round(time.monotonic() - started, 3),
            "cache_predicate_pass": correct_cache, "report_predicate_pass": correct_report,
            "status": "PASS" if passed else "FAIL", "file_sha256_before": initial_hash,
            "file_sha256_after": digest(target), "log_stdout": rel(CONTROLS / (kind + ".stdout.txt")),
            "log_stderr": rel(CONTROLS / (kind + ".stderr.txt")),
            "isolated_profile": {"USERPROFILE": str(PROFILE), "APPDATA": str(PROFILE / ".config"),
                                 "LOCALAPPDATA": str(PROFILE / ".local"), "HOME_modified": False}}


def main():
    for path in (RESOURCES, RUNTIME, PROFILE, CONTROLS, REPORT):
        assert path.resolve().is_relative_to(ROOT.resolve())
    if REPORT.exists() or RUNTIME.exists() or CONTROLS.exists():
        raise SystemExit("Refusing to overwrite previous installation/control/report; preserve prior attempts")
    report = {"status": "IN_PROGRESS", "model_calls": 0, "credentials_read": False,
              "system_installation": False, "method": "Official MSI verified, Windows Installer /a source-image extraction only",
              "author_recalc_script": rel(RECALC), "author_recalc_sha256_before": digest(RECALC),
              "author_source_modified": False,
              "package": {"public_url": URL, "expected_bytes": SIZE, "expected_sha256": HASH,
                          "metadata_source": "https://download.documentfoundation.org/libreoffice/stable/26.8.0/win/x86_64/LibreOffice_26.8.0_Win_x86-64.msi.mirrorlist"},
              "official_references": {
                  "msiexec_admin": "https://learn.microsoft.com/en-us/windows/win32/msi/administrative-installation",
                  "targetdir": "https://learn.microsoft.com/en-us/windows/win32/msi/targetdir",
                  "cli_profile": "https://help.libreoffice.org/latest/en-US/text/shared/guide/start_parameters.html",
                  "portable_distributor": "https://www.libreoffice.org/download-other/"},
              "initial_local_checks": {"soffice_on_path": False, "six_common_binary_paths_present": False,
                  "project_matches": 0, "project_unreachable_paths": 1,
                  "msiexec": "C:/WINDOWS/system32/msiexec.exe", "sevenzip_lessmsi_found": False,
                  "system_vcruntime140_and_msvcp140_exist": True, "openpyxl_actual_import": "3.1.5",
                  "D_free_bytes_observed": 347072610304},
              "source_limitations": ["Original Windows recalc uses ~/.config/libreoffice/4 macro path and no subprocess timeout for final recalc.",
                                     "No binary install or macro presence alone is considered formula-cache verification.",
                                     "MSI /a creates a source image; executable usability is tested here, never inferred from extraction alone.",
                                     "Official parallel-install Wiki and PortableApps package page were inaccessible; no unverified portable package size used."]}
    try:
        package, report["download"] = download()
        RUNTIME.mkdir(parents=True)
        CONTROLS.mkdir(parents=True)
        PROFILE.mkdir(parents=True)
        (PROFILE / ".config").mkdir()
        (PROFILE / ".local").mkdir()
        msi_log = PREFLIGHT / "libreoffice_admin_extract.log"
        assert not msi_log.exists()
        command = ["C:\\WINDOWS\\system32\\msiexec.exe", "/a", str(package), "/qn", "/norestart",
                   "TARGETDIR=" + str(RUNTIME), "/l*v", str(msi_log)]
        print("Extracting official MSI to project-only runtime with /a", flush=True)
        completed = subprocess.run(command, capture_output=True, text=True, timeout=600)
        report["extraction"] = {"method": "msiexec /a /qn /norestart TARGETDIR=<project runtime>",
                                "exit_code": completed.returncode, "target": rel(RUNTIME), "log": rel(msi_log)}
        if completed.returncode != 0:
            raise RuntimeError("Administrative source-image extraction failed with exit " + str(completed.returncode))
        binaries = list(RUNTIME.rglob("soffice.exe"))
        if len(binaries) != 1:
            raise RuntimeError("Expected one extracted soffice.exe, found " + str(len(binaries)))
        executable = binaries[0]
        report["extracted_binary"] = {"path": rel(executable), "sha256": digest(executable),
                                      "bytes": executable.stat().st_size}
        report["controls"] = [run_control("positive", "=200000/12", Decimal("200000") / 12, executable.parent),
                              run_control("division_by_zero", "=1/0", "#DIV/0!", executable.parent)]
        report["status"] = "READY_ORIGINAL_RECALC_REAL_CACHE_POSITIVE_AND_ERROR_CONTROLS_PASS" if all(c["status"] == "PASS" for c in report["controls"]) else "NOT_READY_ORIGINAL_RECALC_CONTROL_FAILURE_RETAINED"
        macro = PROFILE / ".config/libreoffice/4/user/basic/Standard/Module1.xba"
        report["macro_at_original_script_expected_path"] = {"path": rel(macro), "exists": macro.exists(),
                                                           "sha256": digest(macro) if macro.exists() else None}
    except Exception as error:
        report["status"] = "NOT_READY_RUNTIME_PREPARATION_FAILURE_RETAINED"
        report["failure"] = {"type": type(error).__name__, "message": str(error), "traceback": traceback.format_exc()}
    report["author_recalc_sha256_after"] = digest(RECALC)
    assert report["author_recalc_sha256_before"] == report["author_recalc_sha256_after"]
    write(REPORT, report)
    print(json.dumps({"status": report["status"], "report": rel(REPORT), "model_calls": 0}, ensure_ascii=False), flush=True)
    if not report["status"].startswith("READY_"):
        raise SystemExit(2)


if __name__ == "__main__":
    main()
