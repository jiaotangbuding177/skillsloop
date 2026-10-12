"""Native evidence extraction, document rendering and explicit patch application.

No model, risk prediction, reference answer or private business constants.
"""
import argparse, hashlib, json
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn

def load(path):return json.loads(Path(path).read_text('utf-8-sig'))
def save(path,value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

def nodes(document):
    # Keep physical paragraph and table coordinates, including blank paragraphs.
    result=[(f'p{i+1:04}',p) for i,p in enumerate(document.paragraphs)]
    seen={p._p for _,p in result}
    for t,table in enumerate(document.tables):
        for r,row in enumerate(table.rows):
            for c,cell in enumerate(row.cells):
                for i,p in enumerate(cell.paragraphs):
                    # A merged cell occurs under several grid coordinates but
                    # owns one physical paragraph. Index its first coordinate.
                    if p._p not in seen:
                        result.append((f't{t+1}r{r+1}c{c+1}p{i+1}',p));seen.add(p._p)
    return result

def extract(path):
    path=Path(path);kind=path.suffix.lower()
    if kind=='.docx':values=[dict(locator=k,text=p.text) for k,p in nodes(Document(path))]
    elif kind=='.pdf':
        from pypdf import PdfReader
        values=[dict(locator=f'page{i+1}',text=p.extract_text() or '') for i,p in enumerate(PdfReader(path).pages)]
    elif kind in ('.txt','.md'):
        values=[dict(locator=f'line{i+1}',text=line) for i,line in enumerate(path.read_text('utf-8-sig').splitlines())]
    else:raise ValueError('Convert legacy DOC with LibreOffice first; accepted DOCX/PDF/TXT/MD')
    return dict(source=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),records=values,
        text_extraction_only=True,ocr_required_if_empty=not any(r['text'].strip() for r in values))

def style(document):
    normal=document.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'WenQuanYi Zen Hei')
    normal.paragraph_format.space_after=Pt(6)
    for section in document.sections:
        section.top_margin=section.bottom_margin=Cm(2)

def render(data,out):
    if not data.get('title') or not any(data.get(k) for k in ('sections','tables')):
        raise ValueError('A title and substantive sections/tables are required')
    d=Document();style(d);d.add_heading(data['title'],0);md=['# '+data['title'],'']
    if data.get('provenance'):
        d.add_paragraph(data['provenance']);md += [data['provenance'],'']
    for section in data.get('sections',[]):
        if section.get('heading'):d.add_heading(section['heading'],1);md += ['## '+section['heading'],'']
        for text in section.get('paragraphs',[]):
            if not isinstance(text,str):raise ValueError('Paragraph text must be a string')
            d.add_paragraph(text);md += [text,'']
        for point in section.get('bullets',[]):d.add_paragraph(point,style='List Bullet');md += ['- '+point]
        md.append('')
    for table in data.get('tables',[]):
        columns=table['columns'];items=table['rows']
        if not columns or any(len(row)!=len(columns) for row in items):raise ValueError('Table shape mismatch')
        if table.get('heading'):d.add_heading(table['heading'],1)
        tab=d.add_table(rows=1,cols=len(columns));tab.style='Table Grid'
        for cell,text in zip(tab.rows[0].cells,columns):cell.text=str(text)
        for row in items:
            for cell,text in zip(tab.add_row().cells,row):cell.text=str(text)
        clean=lambda s:str(s).replace('|','\\|').replace('\n','<br>')
        md += ['| '+' | '.join(map(clean,columns))+' |','| '+' | '.join(['---']*len(columns))+' |']
        md += ['| '+' | '.join(map(clean,row))+' |' for row in items];md.append('')
    out=Path(out);out.parent.mkdir(parents=True,exist_ok=True);d.save(out)
    out.with_suffix('.md').write_text('\n'.join(md),encoding='utf-8')
    return dict(docx=str(out),markdown=str(out.with_suffix('.md')),paragraphs=len(d.paragraphs),tables=len(d.tables))

def revise(template,data,out):
    doc=Document(template);lookup=dict(nodes(doc));xml={k:p._p.xml for k,p in lookup.items()}
    edited=set();patches=data.get('patches',[])
    for patch in patches:
        key=patch['locator']
        if key in edited or key not in lookup:raise ValueError('Unknown or duplicate paragraph locator')
        paragraph=lookup[key]
        if paragraph.text!=patch['original']:raise ValueError('Exact original text mismatch at '+key)
        if not patch.get('reason'):raise ValueError('Each patch must explain its reason')
        replacement=patch['replacement']
        paragraph.clear();run=paragraph.add_run(replacement)
        if patch.get('bold'):run.bold=True;run.font.color.rgb=RGBColor(0,0,0)
        edited.add(key)
    for clause in data.get('append_clauses',[]):
        if not clause.get('text') or not clause.get('reason'):raise ValueError('New clause and reason required')
        run=doc.add_paragraph().add_run(clause['text']);run.bold=True;run.font.color.rgb=RGBColor(0,0,0)
    if not patches and not data.get('append_clauses'):raise ValueError('No explicit changes requested')
    untouched={k:hashlib.sha256(value.encode()).hexdigest() for k,value in xml.items() if k not in edited}
    assert all(lookup[k]._p.xml==xml[k] for k in untouched),'Unexpected change to protected paragraphs'
    out=Path(out);out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
    reopened=dict(nodes(Document(out)))
    assert all(reopened[k]._p.xml==xml[k] for k in untouched),'Protected XML changed on serialization'
    receipt=dict(template_sha256=hashlib.sha256(Path(template).read_bytes()).hexdigest(),
        output_sha256=hashlib.sha256(out.read_bytes()).hexdigest(),edited_locators=sorted(edited),
        untouched_paragraph_xml_sha256=untouched,untouched_paragraphs_preserved=True,
        document=str(out),legal_assessment_performed_by_script=False)
    save(out.with_suffix('.preservation.json'),receipt);return receipt

def review(data,out):
    # Risk/wording decisions come from the actor, never keyword scoring or fixed outputs.
    original=extract(data['contract']);records={r['locator']:r['text'] for r in original['records']}
    findings=data.get('findings',[])
    for f in findings:
        if f.get('original') not in records.get(f['locator'],'') or not f.get('original'):
            raise ValueError('Review must quote actual text at its locator')
        if not all(f.get(k) for k in ('replacement','reason','level','basis')):
            raise ValueError('Finding requires real revision, reason, actor-assigned level and basis')
    payload=dict(title=data.get('title','合同审查报告'),provenance='结构化审查材料；等级与法律判断由执行者根据证据填写。',
        sections=[dict(heading='审查范围',paragraphs=[str(data.get('standpoint','未知立场')),
            '原件SHA256：'+original['sha256'],data.get('scope','仅对当前已提供合同与资料分析')]),
            dict(heading='待核实事项',bullets=data.get('unknowns',[])),
            dict(heading='审查结论',paragraphs=[data.get('conclusion','尚未填写结论；不得据此认定无风险。')])],
        tables=[dict(heading='逐项修订对照',columns=['定位与原文','建议修改','理由、风险与依据'],rows=[
            [f['locator']+'\n'+f['original'],f['replacement'],f['level']+'\n'+f['reason']+'\n'+f['basis']] for f in findings])])
    return render(payload,out)

def risk(data,out):
    evidence=[]
    for entry in data.get('evidence',[]):
        src=extract(entry['path']);texts={r['locator']:r['text'] for r in src['records']}
        if not entry.get('quote') or entry['quote'] not in texts.get(entry['locator'],''):
            raise ValueError('Evidence quote does not match actual source locator')
        evidence.append([entry['id'],entry['path']+' / '+entry['locator'],entry['quote'],src['sha256']])
    identifiers={r[0] for r in evidence}
    if not evidence or len(identifiers)!=len(evidence):raise ValueError('Unique actual evidence entries required')
    issues=data.get('issues',[])
    for issue in issues:
        if not issue.get('evidence_ids') or not set(issue['evidence_ids'])<=identifiers:
            raise ValueError('Every issue must cite actual indexed evidence')
        if not all(issue.get(k) for k in ('issue','analysis','action')):raise ValueError('Incomplete issue')
    return render(dict(title=data.get('title','法律风险材料梳理'),provenance='上下文重建技能；证据核对通过不等于法律结论自动成立。',
        sections=[dict(heading='任务及立场',paragraphs=[data.get('standpoint','未知立场'),data.get('scope','当前已提供材料')]),
            dict(heading='待核实事项',bullets=data.get('unknowns',[]))],tables=[
            dict(heading='证据索引',columns=['证据','来源及定位','原文摘录','文件SHA256'],rows=evidence),
            dict(heading='问题与建议',columns=['问题','引用证据','分析与依据','建议动作'],rows=[
                [i['issue'],', '.join(i['evidence_ids']),i['analysis'],i['action']] for i in issues])]),out)

def main():
    p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='command',required=True)
    e=sub.add_parser('extract');e.add_argument('source');e.add_argument('--out',required=True)
    for name in ('render','review','risk'):
        s=sub.add_parser(name);s.add_argument('input');s.add_argument('--out',required=True)
    s=sub.add_parser('revise');s.add_argument('template');s.add_argument('input');s.add_argument('--out',required=True)
    a=p.parse_args()
    if a.command=='extract':result=extract(a.source);save(a.out,result)
    elif a.command=='revise':result=revise(a.template,load(a.input),a.out)
    else:result=globals()[a.command](load(a.input),a.out)
    print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':main()
