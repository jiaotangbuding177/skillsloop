from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any


def clean(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(pdf_path: Path) -> dict[str, Any]:
    if not pdf_path.exists():
        return {"error": f"paper file not found: {pdf_path}"}
    try:
        import pymupdf
    except Exception as exc:
        return {"error": f"PyMuPDF/pymupdf unavailable: {exc}"}

    doc = pymupdf.open(pdf_path)
    pages = []
    for i, page in enumerate(doc, start=1):
        pages.append((i, clean(page.get_text("text"))))

    first = pages[0][1] if pages else ""
    keywords = [
        "Fig.",
        "Figure",
        "Supplementary",
        "Data availability",
        "Statistical analysis",
        "Methods",
        "Results",
    ]
    hits = []
    for page_num, text in pages:
        low = text.lower()
        for kw in keywords:
            idx = low.find(kw.lower())
            if idx >= 0:
                hits.append(
                    {
                        "page": page_num,
                        "keyword": kw,
                        "snippet": text[max(0, idx - 300) : min(len(text), idx + 700)],
                    }
                )
    return {
        "pages": len(doc),
        "metadata": doc.metadata,
        "first_page": first[:2500],
        "hits": hits[:40],
    }


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def inspect_xlsx(xlsx_path: Path) -> dict[str, Any]:
    if not xlsx_path.exists():
        return {"error": f"data file not found: {xlsx_path}"}
    try:
        import openpyxl
    except Exception as exc:
        return {"error": f"openpyxl unavailable: {exc}"}

    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        preview = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True):
            preview.append(["" if v is None else str(v) for v in row[: min(ws.max_column, 12)]])
        non_empty_rows = 0
        non_empty_cols = set()
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if any(not is_blank(v) for v in row):
                non_empty_rows += 1
            for col_idx, value in enumerate(row, start=1):
                if not is_blank(value):
                    non_empty_cols.add(col_idx)
        sheets.append(
            {
                "name": ws.title,
                "state": ws.sheet_state,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "non_empty_rows": non_empty_rows,
                "non_empty_cols": len(non_empty_cols),
                "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges][:20],
                "preview": preview,
            }
        )
    return {"sheets": sheets}


def md_escape(value: Any) -> str:
    return str(value).replace("\n", " ").replace("|", "\\|")


def write_report(paper: Path | None, data: Path | None, out: Path) -> None:
    lines = ["# Paper Data Preflight", ""]

    if paper:
        pdf = extract_pdf(paper)
        lines += ["## Paper", "", f"- File: `{paper}`"]
        if "error" in pdf:
            lines.append(f"- Error: {pdf['error']}")
        else:
            lines += [
                f"- Pages: {pdf['pages']}",
                f"- Title metadata: {pdf['metadata'].get('title', '')}",
                f"- DOI/subject metadata: {pdf['metadata'].get('subject', '')}",
                "",
                "### First Page Extract",
                "",
                "```text",
                pdf["first_page"],
                "```",
                "",
                "### Keyword Context",
            ]
            for hit in pdf["hits"][:12]:
                lines += [
                    "",
                    f"- Page {hit['page']} `{hit['keyword']}`: {md_escape(hit['snippet'][:600])}",
                ]

    if data:
        xlsx = inspect_xlsx(data)
        lines += ["", "## Data File", "", f"- File: `{data}`"]
        if "error" in xlsx:
            lines.append(f"- Error: {xlsx['error']}")
        else:
            lines += ["", "| Sheet | Size | Non-empty | Merged ranges |", "| --- | --- | --- | --- |"]
            for sheet in xlsx["sheets"]:
                lines.append(
                    f"| {md_escape(sheet['name'])} | {sheet['max_row']}x{sheet['max_column']} | "
                    f"{sheet['non_empty_rows']} rows, {sheet['non_empty_cols']} cols | "
                    f"{md_escape(', '.join(sheet['merged_ranges'][:5]))} |"
                )
            for sheet in xlsx["sheets"]:
                lines += ["", f"### Preview: {sheet['name']}", "", "```text"]
                for row in sheet["preview"]:
                    lines.append(str(row))
                lines.append("```")

    lines += [
        "",
        "## Manual Review Reminder",
        "",
        "Use this preflight only to orient the review. Final conclusions must compare the paper's expected figures, groups, sample sizes, and claims against the data file.",
    ]
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--paper", type=Path)
    parser.add_argument("--data", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    write_report(args.paper, args.data, args.out)
    print(args.out)


if __name__ == "__main__":
    main()
